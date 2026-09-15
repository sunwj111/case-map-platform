import zipfile
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.produce import parse_cases, parse_knowledge


def test_creates_standard_batch_with_dual_input_gate(produce_service):
    batch = produce_service.create({"domain": " 家装 ", "system": " 报价 ", "mode": "standard"})
    assert batch["mode"] == "standard"
    assert batch["domain"] == "家装"
    assert batch["system"] == "报价"
    assert batch["moduleName"] == ""
    assert batch["casesRequired"] is True
    assert batch["knowledgeRequired"] is True
    assert batch["status"] == "WAITING_FOR_CASES_AND_KNOWLEDGE"
    assert batch["missingInputs"] == ["cases", "knowledge"]
    assert batch["readyForKeywordCalibration"] is False
    assert batch["readyForMapDraft"] is False


def test_advances_only_after_cases_knowledge_and_keyword_confirmation(produce_service):
    batch = produce_service.create({"domain": "家装", "system": "报价"})
    csv_content = (
        "用例编号,用例名称,步骤,预期结果,场景,功能点\n"
        "TC-001,数量价汇总,提交报价,金额正确,金额计算与汇总,数量价汇总\n"
    ).encode("utf-8")
    knowledge = "场景：金额计算与汇总\n功能点：数量价汇总\n节点：报价计算节点"

    produce_service.import_cases(batch["id"], "cases.csv", csv_content)
    after_cases = produce_service.get(batch["id"])
    assert after_cases["status"] == "WAITING_FOR_KNOWLEDGE"
    with pytest.raises(ValueError):
        produce_service.confirm_keywords(batch["id"], None)

    produce_service.import_knowledge(batch["id"], "knowledge.md", knowledge.encode("utf-8"))
    after_knowledge = produce_service.get(batch["id"])
    assert after_knowledge["status"] == "READY_FOR_KEYWORD_CALIBRATION"
    assert after_knowledge["readyForKeywordCalibration"] is True
    with pytest.raises(ValueError):
        produce_service.confirm_keywords(batch["id"], None)

    produce_service.extract_knowledge(batch["id"])
    after_confirmation = produce_service.confirm_keywords(batch["id"], None)
    assert after_confirmation["status"] == "READY_FOR_MAP_DRAFT"
    assert after_confirmation["readyForMapDraft"] is True

    draft = produce_service.generate_draft(batch["id"])
    generated = produce_service.get(batch["id"])
    assert generated["status"] == "REVIEW_IN_PROGRESS"
    assert generated["reviewQueueGenerated"] is True
    assert draft["stats"]["totalCases"] == 1
    assert draft["cases"][0]["feature"] == "数量价汇总"
    reviews = produce_service.query_reviews(batch["id"], None, None, None, None, None)
    assert reviews["total"] == 1


def test_uncovered_rules_and_apis_become_gap_candidates(produce_service):
    batch = produce_service.create({"domain": "家装", "system": "报价"})
    csv_content = (
        "用例编号,用例名称,步骤,预期结果,场景,功能点\n"
        "TC-001,数量价汇总,提交报价,金额正确,金额计算与汇总,数量价汇总\n"
    ).encode("utf-8")
    knowledge = (
        "场景：金额计算与汇总\n"
        "功能点：数量价汇总\n"
        "功能点：去利润价计算\n"
        "规则：金额必须大于零\n"
        "接口：POST /quote/submit\n"
        "POST /quote/audit\n"
    )
    produce_service.import_cases(batch["id"], "cases.csv", csv_content)
    produce_service.import_knowledge(batch["id"], "knowledge.md", knowledge.encode("utf-8"))
    produce_service.extract_knowledge(batch["id"])
    produce_service.confirm_keywords(batch["id"], None)
    draft = produce_service.generate_draft(batch["id"])
    feature_gaps = [item for item in draft["gaps"] if item["kind"] == "功能点缺口"]
    rule_gaps = [item for item in draft["gaps"] if item["kind"] == "规则缺口"]
    api_gaps = [item for item in draft["gaps"] if item["kind"] == "接口缺口"]
    assert any(item["feature"] == "去利润价计算" for item in feature_gaps)
    assert all(item["scene"] == "金额计算与汇总" for item in feature_gaps)
    assert all(item["caseName"].startswith("验证") for item in draft["gaps"])
    assert all("1." in (item.get("step") or "") and "1." in (item.get("expected") or "") for item in draft["gaps"])
    profit_gaps = [item for item in feature_gaps if item["feature"] == "去利润价计算"]
    assert len(profit_gaps) == 9
    assert all("POST /quote/submit" in (item.get("api") or "") for item in profit_gaps)
    assert {item["priority"] for item in profit_gaps} == {"P1", "P2", "P3"}
    assert any("主流程" in item["caseName"] for item in profit_gaps)
    assert any("必填" in item["caseName"] for item in profit_gaps)
    assert any("边界" in item["caseName"] for item in profit_gaps)
    assert any("特殊字符" in item["caseName"] for item in profit_gaps)
    assert any("重复" in item["caseName"] for item in profit_gaps)
    assert any("权限" in item["caseName"] for item in profit_gaps)
    assert any("联动" in item["caseName"] for item in profit_gaps)
    assert any("文案" in item["caseName"] for item in profit_gaps)
    assert any("按钮" in item["caseName"] or "跳转" in item["caseName"] for item in profit_gaps)
    assert any("金额必须大于零" in item["feature"] for item in rule_gaps)
    assert len([item for item in rule_gaps if "金额必须大于零" in item["feature"]]) == 9
    assert any("quote" in item["feature"] and "submit" in item["feature"] for item in api_gaps)
    assert any("主流程" in item["caseName"] for item in api_gaps)
    assert any("必填" in item["caseName"] for item in api_gaps)
    reviews = produce_service.query_reviews(batch["id"], None, None, None, None, None)
    review_kinds = {item["kind"] for item in reviews["items"]}
    assert {"历史用例", "功能点缺口", "规则缺口", "接口缺口"} <= review_kinds
    historical = [item for item in reviews["items"] if item["kind"] == "历史用例"]
    assert historical[0]["api"] == ""
    api_reviews = [item for item in reviews["items"] if item["kind"] == "接口缺口"]
    assert any(item["api"] == "POST /quote/submit" for item in api_reviews)
    assert any(item["api"] == "POST /quote/audit" for item in api_reviews)


def test_gap_scene_follows_knowledge_section(produce_service):
    batch = produce_service.create({"domain": "家装", "system": "报价"})
    csv_content = (
        "用例编号,用例名称,步骤,预期结果,场景,功能点\n"
        "TC-001,数量价汇总,提交报价,金额正确,金额计算与汇总,数量价汇总\n"
    ).encode("utf-8")
    knowledge = (
        "场景：金额计算与汇总\n"
        "功能点：数量价汇总\n"
        "场景：造价提交与审核\n"
        "功能点：审核拒绝回退\n"
        "规则：免审条件命中后可以自动通过\n"
    )
    produce_service.import_cases(batch["id"], "cases.csv", csv_content)
    produce_service.import_knowledge(batch["id"], "knowledge.md", knowledge.encode("utf-8"))
    produce_service.extract_knowledge(batch["id"])
    produce_service.confirm_keywords(batch["id"], None)
    draft = produce_service.generate_draft(batch["id"])
    reject_gaps = [item for item in draft["gaps"] if item["feature"] == "审核拒绝回退"]
    assert reject_gaps
    assert all(item["scene"] == "造价提交与审核" for item in reject_gaps)
    rule_gaps = [item for item in draft["gaps"] if item["kind"] == "规则缺口"]
    assert rule_gaps
    assert all(item["scene"] == "造价提交与审核" for item in rule_gaps)


def test_history_case_binds_api_only_when_text_hits_knowledge(produce_service):
    batch = produce_service.create({"domain": "家装", "system": "报价"})
    csv_content = (
        "用例编号,用例名称,步骤,预期结果,场景,功能点\n"
        "TC-001,提交报价,调用 POST /quote/submit,返回成功,金额计算与汇总,数量价汇总\n"
        "TC-002,数量价汇总,打开汇总页,金额正确,金额计算与汇总,数量价汇总\n"
        "TC-003,其它功能,打开其它页,展示正确,金额计算与汇总,优惠计算\n"
    ).encode("utf-8")
    knowledge = (
        "场景：金额计算与汇总\n"
        "功能点：数量价汇总\n"
        "接口：POST /quote/submit\n"
        "接口：POST /quote/audit\n"
    )
    produce_service.import_cases(batch["id"], "cases.csv", csv_content)
    produce_service.import_knowledge(batch["id"], "knowledge.md", knowledge.encode("utf-8"))
    produce_service.extract_knowledge(batch["id"])
    produce_service.confirm_keywords(batch["id"], None)
    draft = produce_service.generate_draft(batch["id"])
    by_id = {item["originalCaseId"]: item for item in draft["cases"]}
    assert "POST /quote/submit" in by_id["TC-001"]["api"]
    assert "POST /quote/submit" in by_id["TC-002"]["api"]
    assert by_id["TC-003"]["api"] == ""
    reviews = produce_service.query_reviews(batch["id"], None, None, None, None, None)
    first_review = next(item for item in reviews["items"] if item["originalCaseId"] == "TC-001")
    produce_service.update_review(first_review["id"], {"api": "POST /quote/submit; POST /quote/audit"})
    updated = produce_service.query_reviews(batch["id"], None, None, None, None, None)
    edited = next(item for item in updated["items"] if item["originalCaseId"] == "TC-001")
    assert edited["api"] == "POST /quote/submit; POST /quote/audit"


def test_outline_headings_not_used_as_scene_or_feature(produce_service):
    batch = produce_service.create({
        "domain": "家装",
        "system": "报价",
        "processNode": "造价审核",
    })
    csv_content = (
        "用例编号,用例名称,步骤,预期结果,场景,功能点\n"
        "TC-001,历史报价,提交,成功,金额计算与汇总,数量价汇总\n"
    ).encode("utf-8")
    knowledge = (
        "# 一、场景（业务背景）\n"
        "造价审核背景说明\n"
        "# 二、功能点\n"
        "- 造价审核\n"
    )
    produce_service.import_cases(batch["id"], "cases.csv", csv_content)
    produce_service.import_knowledge(batch["id"], "knowledge.md", knowledge.encode("utf-8"))
    produce_service.extract_knowledge(batch["id"])
    keywords = produce_service.get(batch["id"])["keywordExtraction"]
    feature_texts = [item["text"] for item in keywords["features"]]
    scene_texts = [item["text"] for item in keywords["scenes"]]
    assert "二、功能点" not in feature_texts
    assert "一、（业务背景）" not in scene_texts
    assert "造价审核" in feature_texts
    produce_service.confirm_keywords(batch["id"], None)
    draft = produce_service.generate_draft(batch["id"])
    feature_gaps = [item for item in draft["gaps"] if item["kind"] == "功能点缺口"]
    assert feature_gaps
    assert all(item["scene"] == "金额计算与汇总" for item in feature_gaps)
    assert all(item["feature"] == "造价审核" for item in feature_gaps)
    assert all("二、功能点" not in item["caseName"] for item in feature_gaps)
    assert all("业务背景" not in (item["scene"] or "") for item in feature_gaps)


def test_gap_scene_falls_back_to_history_not_process_node(produce_service):
    batch = produce_service.create({
        "domain": "家装",
        "system": "报价",
        "processNode": "金额计算",
        "moduleName": "报价计算",
    })
    csv_content = (
        "用例编号,用例名称,步骤,预期结果,场景,功能点\n"
        "TC-001,无权限访问,打开规则,提示无权限,权限校验,状态枚举\n"
    ).encode("utf-8")
    knowledge = (
        "功能点：审核拒绝回退\n"
        "规则：无权限不可操作\n"
        "接口：POST /check-status\n"
    )
    produce_service.import_cases(batch["id"], "cases.csv", csv_content)
    produce_service.import_knowledge(batch["id"], "knowledge.md", knowledge.encode("utf-8"))
    produce_service.extract_knowledge(batch["id"])
    produce_service.confirm_keywords(batch["id"], None)
    draft = produce_service.generate_draft(batch["id"])
    assert draft["gaps"]
    assert all(item["scene"] == "权限校验" for item in draft["gaps"])
    assert all(item["scene"] != "金额计算" for item in draft["gaps"])
    keywords = produce_service.get(batch["id"])["keywordExtraction"]
    assert all(item["text"] != "金额计算" for item in keywords.get("scenes") or [])


def test_feature_cases_bind_section_api(produce_service):
    batch = produce_service.create({"domain": "家装", "system": "报价"})
    csv_content = (
        "用例编号,用例名称,步骤,预期结果,场景,功能点\n"
        "TC-001,无权限访问,打开规则页,提示无权限,权限校验,状态枚举\n"
    ).encode("utf-8")
    knowledge = (
        "场景：权限校验\n"
        "功能点：状态枚举\n"
        "接口：POST /check-status\n"
        "功能点：审核拒绝回退\n"
        "接口：POST /audit/reject\n"
    )
    produce_service.import_cases(batch["id"], "cases.csv", csv_content)
    produce_service.import_knowledge(batch["id"], "knowledge.md", knowledge.encode("utf-8"))
    produce_service.extract_knowledge(batch["id"])
    produce_service.confirm_keywords(batch["id"], None)
    draft = produce_service.generate_draft(batch["id"])
    historical = draft["cases"][0]
    assert historical["api"] == "POST /check-status"
    reject_gaps = [item for item in draft["gaps"] if item["feature"] == "审核拒绝回退"]
    assert reject_gaps
    assert all(item["api"] == "POST /audit/reject" for item in reject_gaps)


def test_markdown_table_knowledge_binds_feature_apis(produce_service):
    batch = produce_service.create({"domain": "家装", "system": "报价"})
    csv_content = (
        "用例编号,用例名称,步骤,预期结果,场景,功能点\n"
        "TC-001,接单成功,点击接单,状态变为已接单,造价单审核,接单\n"
    ).encode("utf-8")
    knowledge = """## 二、功能点
| # | 功能点 | 实现位置 | 说明 |
|---|---|---|---|
| 4 | **接单**（认领） | `DesignCostServiceImpl#accept` | 把审核人置为当前用户 |
| 5 | **转派** | `DesignCostServiceImpl#transform` | 批量改审核人 |
| 6 | **催审** | `CostAuditRpcServiceImpl#urgentReview` | 写催审时间 |

## 五、接口
| 方法 | 路径 | 入参 |
|---|---|---|
| POST | `/costAudit/findList` | `CostAuditListReq` |
| POST | `/costAudit/accept` | `CostAuditAcceptReq` |
| POST | `/costAudit/transform` | `CostAuditTransformReq` |
| POST | `/costAudit/urgentReview` | `UrgentReviewReq` |
"""
    produce_service.import_cases(batch["id"], "cases.csv", csv_content)
    produce_service.import_knowledge(batch["id"], "knowledge.md", knowledge.encode("utf-8"))
    produce_service.extract_knowledge(batch["id"])
    keywords = produce_service.get(batch["id"])["keywordExtraction"]
    api_texts = [item["text"] for item in keywords["apis"]]
    assert "POST /costAudit/accept" in api_texts
    accept = next(item for item in keywords["apis"] if "accept" in item["text"])
    assert accept.get("relatedFeature") == "接单"
    produce_service.confirm_keywords(batch["id"], {
        "apis": [item["text"] for item in keywords["apis"]],
        "features": [item["text"] for item in keywords["features"]],
    })
    draft = produce_service.generate_draft(batch["id"])
    assert draft["cases"][0]["api"] == "POST /costAudit/accept"
    transfer_gaps = [item for item in draft["gaps"] if item["feature"] == "转派" and item["kind"] == "功能点缺口"]
    assert transfer_gaps
    assert all("POST /costAudit/transform" in (item.get("api") or "") for item in transfer_gaps)
    urgent_gaps = [item for item in draft["gaps"] if item["feature"] == "催审" and item["kind"] == "功能点缺口"]
    assert urgent_gaps
    assert all("POST /costAudit/urgentReview" in (item.get("api") or "") for item in urgent_gaps)


def test_audit_history_cases_bind_check_design_cost(produce_service):
    batch = produce_service.create({"domain": "家装", "system": "报价"})
    csv_content = (
        "用例编号,用例名称,步骤,预期结果,场景,功能点\n"
        "TC-001,造价审核专员成功审核造价单,点击审核通过,审核通过,造价单审核,人工审核\n"
        "TC-002,设计师成功提交造价审核,超级ZD设计师提交造价审核,提交成功,造价单审核,人工审核\n"
        "TC-003,造价单自动审核通过,提交造价单,自动通过,造价单审核,自动审核通过\n"
    ).encode("utf-8")
    knowledge = (
        "场景：造价单审核\n"
        "功能点：审核通过/驳回\n"
        "功能点：自动审核通过\n"
        "CostAuditApi /checkDesignCost\n"
        "提交(submitDesignCost)\n"
        "whetherAutoSubmit\n"
        "接口：POST /costAudit/findList\n"
    )
    produce_service.import_cases(batch["id"], "cases.csv", csv_content)
    produce_service.import_knowledge(batch["id"], "knowledge.md", knowledge.encode("utf-8"))
    produce_service.extract_knowledge(batch["id"])
    produce_service.confirm_keywords(batch["id"], None)
    keywords = produce_service.get(batch["id"])["keywordExtraction"]
    api_texts = [item["text"] for item in keywords["apis"]]
    assert any("checkDesignCost" in text for text in api_texts)
    draft = produce_service.generate_draft(batch["id"])
    by_id = {item["originalCaseId"]: item for item in draft["cases"]}
    assert "checkDesignCost" in (by_id["TC-001"]["api"] or "")
    assert "submitDesignCost" in (by_id["TC-002"]["api"] or "")
    assert "whetherAutoSubmit" in (by_id["TC-003"]["api"] or "") or "autoSubmit" in (by_id["TC-003"]["api"] or "")
    pass_gaps = [item for item in draft["gaps"] if item["feature"] == "审核通过/驳回"]
    assert pass_gaps
    assert all("checkDesignCost" in (item.get("api") or "") for item in pass_gaps)


def test_pasted_knowledge_extends_uploaded_file(produce_service):
    batch = produce_service.create({"domain": "家装", "system": "报价"})
    produce_service.import_knowledge(batch["id"], "uploaded.md", "# 功能点\n- 自动审核判定\n".encode("utf-8"))
    produce_service.import_knowledge_text(batch["id"], "人工补充", "规则：金额必须大于零")
    updated = produce_service.get(batch["id"])
    assert "自动审核判定" in updated["knowledgeImport"]["text"]
    assert "金额必须大于零" in updated["knowledgeImport"]["text"]
    assert updated["knowledgeImport"]["metadata"]["format"] == "combined-knowledge"
    assert updated["keywordExtraction"]["features"]
    assert updated["keywordExtraction"]["rules"]


def test_rejects_unsupported_mode_and_unknown_target(produce_service):
    with pytest.raises(ValueError):
        produce_service.create({"domain": "家装", "system": "报价", "mode": "casesOnly"})
    with pytest.raises(ValueError, match="系统不能包含斜杠或换行"):
        produce_service.create({"domain": "家装", "system": "https://wiki.example.com/sys"})


def test_creates_batch_from_platform_project_fields(produce_service):
    batch = produce_service.create({
        "ziroom_domain": "收房",
        "system_ref": "租住系统",
        "module_name": "签约",
    })
    assert batch["domain"] == "收房"
    assert batch["system"] == "租住系统"
    assert batch["moduleName"] == "签约"
    assert batch["ziroom_domain"] == "收房"
    assert batch["system_ref"] == "租住系统"
    assert batch["module_name"] == "签约"


def test_rejects_mismatched_platform_aliases(produce_service):
    with pytest.raises(ValueError, match="领域与平台字段不一致"):
        produce_service.create({
            "domain": "家装",
            "ziroom_domain": "收房",
            "system": "报价",
        })


def test_platform_fields_do_not_change_cleaning_targets(produce_service):
    batch = produce_service.create({
        "ziroom_domain": "家装",
        "system_ref": "报价",
        "module_name": "报价计算",
    })
    csv_content = (
        "用例编号,用例名称,步骤,预期结果,场景,功能点\n"
        "TC-001,数量价汇总,提交报价,金额正确,金额计算与汇总,数量价汇总\n"
    ).encode("utf-8")
    produce_service.import_cases(batch["id"], "cases.csv", csv_content)
    produce_service.import_knowledge_text(
        batch["id"],
        "粘贴",
        "场景：金额计算与汇总\n功能点：数量价汇总\n节点：报价计算节点\n",
    )
    produce_service.confirm_keywords(batch["id"], None)
    draft = produce_service.generate_draft(batch["id"])
    cleaned = draft["cases"][0]
    assert cleaned["scene"] == "金额计算与汇总"
    assert cleaned["feature"] == "数量价汇总"
    assert "testPoint" not in cleaned
    assert "flowNode" in cleaned


def test_creates_batch_with_module_name(produce_service):
    batch = produce_service.create({
        "domain": "家装",
        "system": "报价",
        "moduleName": " 报价计算 ",
    })
    assert batch["moduleName"] == "报价计算"


def test_rejects_module_name_with_slash(produce_service):
    with pytest.raises(ValueError, match="模块不能包含斜杠或换行"):
        produce_service.create({"domain": "家装", "system": "报价", "moduleName": "报价/计算"})


def test_parses_csv_with_fuzzy_headers():
    csv_content = (
        '用例编号,用例名称,操作步骤,预期结果,业务场景,功能点,所属模块\n'
        'TC-001,"数量,价格汇总","输入数量\n并提交",金额正确,金额计算与汇总,数量价汇总,报价\n'
    ).encode("utf-8")
    result = parse_cases("cases.csv", csv_content)
    assert len(result["rows"]) == 1
    assert result["rows"][0]["originalCaseId"] == "TC-001"
    assert result["rows"][0]["caseName"] == "数量,价格汇总"
    assert result["rows"][0]["step"] == "输入数量 并提交"
    assert result["rows"][0]["feature"] == "数量价汇总"
    assert len(result["metadata"]["sha256"]) == 64


def test_reports_missing_columns_and_generates_temporary_case_id():
    csv_content = "用例名称,预期结果\n固定价汇总,金额正确\n".encode("utf-8")
    result = parse_cases("cases.csv", csv_content)
    assert result["rows"][0]["originalCaseId"] == "IMPORT-ROW-2"
    codes = {risk["code"] for risk in result["risks"]}
    assert "MISSING_CASE_ID_COLUMN" in codes
    assert "MISSING_STEP_COLUMN" in codes
    assert "MISSING_ORIGINAL_CASE_ID" in codes


def test_parses_first_excel_sheet():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "历史用例"
    sheet.append(["CaseName", "Step", "Expected"])
    sheet.append(["优惠后汇总", "提交报价", "结果正确"])
    workbook.create_sheet("忽略工作表")
    buffer = BytesIO()
    workbook.save(buffer)
    result = parse_cases("cases.xlsx", buffer.getvalue())
    assert len(result["rows"]) == 1
    assert result["rows"][0]["caseName"] == "优惠后汇总"
    assert result["metadata"]["sheetName"] == "历史用例"
    assert result["metadata"]["sheetCount"] == 2


def test_parses_xmind_as_historical_cases():
    content_json = """[{
      "rootTopic": {
        "title": "金额计算",
        "children": {"attached": [{
          "title": "数量价汇总",
          "children": {"attached": [{"title": "正常汇总"}]}
        }]}
      }
    }]"""
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("content.json", content_json)
    result = parse_cases("cases.xmind", buffer.getvalue())
    assert result["metadata"]["format"] == "xmind-json"
    assert len(result["rows"]) == 1
    assert result["rows"][0]["caseName"] == "正常汇总"
    assert any(risk["code"] == "GENERATED_XMIND_CASE_IDS" for risk in result["risks"])


def test_parse_knowledge_excel():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["场景：金额计算与汇总", "功能点：数量价汇总"])
    buffer = BytesIO()
    workbook.save(buffer)
    result = parse_knowledge("knowledge.xlsx", buffer.getvalue())
    assert "数量价汇总" in result["text"]
    assert result["text"].startswith("# 从 Excel 导入的知识草稿")


def test_review_confirm_and_publish(produce_service, case_store):
    batch = produce_service.create({"domain": "家装", "system": "报价"})
    csv_content = (
        "用例编号,用例名称,步骤,预期结果,场景,功能点\n"
        "TC-001,数量价汇总,提交报价,金额正确,金额计算与汇总,数量价汇总\n"
    ).encode("utf-8")
    produce_service.import_cases(batch["id"], "cases.csv", csv_content)
    produce_service.import_knowledge_text(
        batch["id"],
        "粘贴",
        "场景：金额计算与汇总\n功能点：数量价汇总\n",
    )
    produce_service.confirm_keywords(batch["id"], None)
    produce_service.generate_draft(batch["id"])
    reviews = produce_service.query_reviews(batch["id"], None, "待评审", None, None, None)
    review_id = reviews["items"][0]["id"]
    produce_service.confirm_review(review_id, {"operator": "测试员"})
    published = produce_service.publish_batch(batch["id"], {"operator": "测试员"})
    assert published["publishedCount"] == 1
    official = case_store.items[published["caseIds"][0]]
    assert official["lifecycle"] == "已发布"
    assert official["moduleName"] == ""
    assert official["featureKey"] == "家装/报价/金额计算与汇总/数量价汇总"
    assert official["api"] == ""


def test_publish_copies_module_name_without_changing_feature_key(produce_service, case_store):
    batch = produce_service.create({"domain": "家装", "system": "报价", "moduleName": "报价计算"})
    csv_content = (
        "用例编号,用例名称,步骤,预期结果,场景,功能点\n"
        "TC-001,数量价汇总,提交报价,金额正确,金额计算与汇总,数量价汇总\n"
    ).encode("utf-8")
    produce_service.import_cases(batch["id"], "cases.csv", csv_content)
    produce_service.import_knowledge_text(
        batch["id"],
        "粘贴",
        "场景：金额计算与汇总\n功能点：数量价汇总\n",
    )
    produce_service.confirm_keywords(batch["id"], None)
    produce_service.generate_draft(batch["id"])
    reviews = produce_service.query_reviews(batch["id"], None, "待评审", None, None, None)
    produce_service.confirm_review(reviews["items"][0]["id"], {"operator": "测试员"})
    published = produce_service.publish_batch(batch["id"], {"operator": "测试员"})
    official = case_store.items[published["caseIds"][0]]
    assert official["moduleName"] == "报价计算"
    assert official["featureKey"] == "家装/报价/金额计算与汇总/数量价汇总"


def _draft_cleaned_case(produce_service, csv_line: str, knowledge: str) -> dict:
    batch = produce_service.create({"domain": "家装", "system": "报价"})
    csv_content = (
        "用例编号,用例名称,步骤,预期结果,场景,功能点\n" + csv_line
    ).encode("utf-8")
    produce_service.import_cases(batch["id"], "cases.csv", csv_content)
    produce_service.import_knowledge_text(batch["id"], "粘贴", knowledge)
    produce_service.extract_knowledge(batch["id"])
    produce_service.confirm_keywords(batch["id"], None)
    return produce_service.generate_draft(batch["id"])["cases"][0]


def test_unrecognized_scene_falls_back_to_history_scene(produce_service):
    cleaned = _draft_cleaned_case(
        produce_service,
        "TC-001,无权限访问,尝试打开规则,提示无权限,--,DesignCheckStatusEnum\n",
        "规则：无权限不可操作\n接口：POST /check-status\n",
    )
    assert cleaned["scene"] == "未识别场景"
    assert cleaned["feature"] == "DesignCheckStatusEnum"
    assert all("历史用例场景兜底" not in item for item in cleaned["matchEvidence"])

    cleaned = _draft_cleaned_case(
        produce_service,
        "TC-002,无权限访问,尝试打开规则,提示无权限,权限校验,DesignCheckStatusEnum\n",
        "规则：无权限不可操作\n接口：POST /check-status\n",
    )
    assert cleaned["scene"] == "权限校验"
    assert "知识未命中场景，已用历史用例场景兜底" in cleaned["matchEvidence"]


def test_knowledge_scene_wins_over_history_placeholder(produce_service):
    cleaned = _draft_cleaned_case(
        produce_service,
        "TC-001,打开金额计算与汇总,提交报价,金额正确,--,数量价汇总\n",
        "场景：金额计算与汇总\n功能点：数量价汇总\n",
    )
    assert cleaned["scene"] == "金额计算与汇总"
    assert "命中知识场景：金额计算与汇总" in cleaned["matchEvidence"]
    assert "历史用例场景兜底" not in " ".join(cleaned["matchEvidence"])
