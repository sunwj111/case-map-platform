from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from threading import Lock
from typing import Any
from uuid import uuid4

from app.assets import CaseAssetStore
from app.feature_key import join
from app.hierarchy import HierarchyStore
from app.platform_fields import resolve_project_context
from app.xmind import parse_xmind_archive

CASE_ID_KEYS = ["caseid", "用例id", "用例编号", "编号", "id"]
CASE_NAME_KEYS = ["casename", "用例名称", "用例名", "name", "title", "标题", "用例"]
STEP_KEYS = ["step", "steps", "步骤", "操作步骤", "测试步骤"]
EXPECTED_KEYS = ["expectedresults", "expected", "预期结果", "期望结果", "预期", "期望"]
SCENE_KEYS = ["scene", "场景", "业务场景", "流程节点", "节点"]
FEATURE_KEYS = ["functionpoint", "feature", "功能点", "featurename", "功能", "模块功能"]
MODULE_KEYS = ["module", "belongplatform", "模块", "平台", "所属模块", "belongmodule"]
STOP_WORDS = {
    "以及", "或者", "并且", "然后", "进行", "可以", "需要", "如果", "当", "的", "了", "是", "在",
    "我们", "系统", "用户", "测试", "用例", "知识库", "流程", "业务", "说明", "如下", "包括",
    "主要", "相关", "一般", "时候", "之后", "之前", "通过", "完成", "支持",
}


def _now() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def sha256_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def extension(file_name: str | None) -> str:
    name = "" if file_name is None else file_name.strip().lower()
    index = name.rfind(".")
    return "" if index < 0 else name[index + 1:]


def decode_text(content: bytes) -> tuple[str, dict[str, Any] | None]:
    try:
        text = content.decode("utf-8")
        risk = None
    except UnicodeDecodeError:
        text = content.decode("gb18030", errors="replace")
        risk = {"code": "ENCODING_FALLBACK", "message": "文件不是有效 UTF-8，已按 GB18030 解析", "row": None}
    if text.startswith("\ufeff"):
        text = text[1:]
    return text, risk


def _compact(value: Any) -> str:
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip()


_PLACEHOLDER_LABELS = {
    "",
    "-",
    "--",
    "---",
    "—",
    "–",
    "/",
    "无",
    "空",
    "无场景",
    "暂无",
    "未识别场景",
    "待确认场景",
    "未识别功能点",
    "null",
    "none",
    "n/a",
    "na",
}


def _usable_label(value: Any) -> str:
    text = _compact(value)
    if not text:
        return ""
    if text in _PLACEHOLDER_LABELS or text.lower() in _PLACEHOLDER_LABELS:
        return ""
    if set(text) <= {"-", "—", "–", "_", ".", "/", "\\"}:
        return ""
    return text


_OUTLINE_PREFIX = re.compile(r"^[（(]*[一二三四五六七八九十百零〇0-9]+[、.．)）]\s*")
_GENERIC_SECTION_LABELS = {
    "功能点", "功能", "能力", "场景", "业务场景", "业务背景", "背景", "概述", "简介",
    "规则", "校验", "约束", "接口", "API", "api", "节点", "流程节点", "目录", "说明",
}


def _strip_outline_prefix(text: str) -> str:
    return _OUTLINE_PREFIX.sub("", (text or "").strip()).strip()


def _parenthetical_name(text: str) -> str:
    matched = re.search(r"[（(]([^）)]+)[）)]\s*$", text or "")
    if not matched:
        return ""
    return matched.group(1).strip()


def _is_generic_section_label(text: str) -> bool:
    compact = _compact(text)
    if not compact:
        return True
    inner = _parenthetical_name(compact)
    if inner and inner not in _GENERIC_SECTION_LABELS:
        return False
    stripped = _strip_outline_prefix(compact)
    stripped = re.sub(r"[（(][^）)]*[）)]", "", stripped).strip("、.．:： ")
    stripped = stripped.replace("业务场景", "").replace("功能点", "").strip()
    if stripped.endswith("场景") and stripped != "场景":
        stripped = stripped[:-2].strip()
    return (not stripped) or stripped in _GENERIC_SECTION_LABELS


def _is_outline_label(text: str) -> bool:
    return _is_generic_section_label(text)


def _is_section_heading_line(title: str) -> bool:
    stripped = _strip_outline_prefix(title)
    return bool(re.match(
        r"^(?:（[^）]*）)?(?:业务)?(?:场景|功能点|功能|能力|规则|校验|接口|API|节点)",
        stripped,
    ))


def _business_label(text: str) -> str:
    compact = _compact(text)
    if not compact:
        return ""
    inner = _parenthetical_name(compact)
    if inner and not _is_generic_section_label(inner):
        return _compact(inner)
    stripped = _strip_outline_prefix(compact)
    stripped = re.sub(r"[（(][^）)]*[）)]", "", stripped).strip("、.．:： ")
    stripped = stripped.replace("业务场景", "").replace("功能点", "").strip()
    if "场景" in stripped and stripped != "场景":
        stripped = stripped.replace("场景", "").strip()
    if _is_generic_section_label(stripped):
        return ""
    return stripped


def _normalize_header(value: str) -> str:
    return re.sub(r"[\s_\-]+", "", value).lower()


def _find_header_index(headers: list[str], keys: list[str]) -> int:
    normalized_keys = [_normalize_header(key) for key in keys]
    for index, header in enumerate(headers):
        normalized = _normalize_header(header)
        if normalized in normalized_keys:
            return index
    for index, header in enumerate(headers):
        normalized = _normalize_header(header)
        if len(normalized) < 2:
            continue
        if any(len(key) >= 2 and (key in normalized or normalized in key) for key in normalized_keys):
            return index
    return -1


class JsonStore:
    def __init__(self, data_file: Path, collection_key: str) -> None:
        self.data_file = data_file
        self.collection_key = collection_key
        self._lock = Lock()
        self.items: dict[str, dict[str, Any]] = {}
        if data_file.exists():
            payload = json.loads(data_file.read_text(encoding="utf-8"))
            for item in payload.get(self.collection_key) or []:
                self.items[item["id"]] = item

    def save(self) -> None:
        self.data_file.parent.mkdir(parents=True, exist_ok=True)
        payload = {"updatedAt": _now(), self.collection_key: list(self.items.values())}
        temporary = self.data_file.with_suffix(self.data_file.suffix + ".tmp")
        temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(self.data_file)


def refresh_gate(batch: dict[str, Any]) -> None:
    batch["casesRequired"] = True
    batch["knowledgeRequired"] = batch.get("mode") == "standard"
    batch["readyForKeywordCalibration"] = bool(batch.get("casesImported")) and (
        not batch["knowledgeRequired"] or bool(batch.get("knowledgeImported"))
    )
    batch["readyForMapDraft"] = batch["readyForKeywordCalibration"] and bool(batch.get("keywordsConfirmed"))
    missing = []
    if not batch.get("casesImported"):
        missing.append("cases")
    if batch["knowledgeRequired"] and not batch.get("knowledgeImported"):
        missing.append("knowledge")
    batch["missingInputs"] = missing
    published_count = int(batch.get("publishedCount") or 0)
    if published_count > 0:
        batch["status"] = "ASSETS_PUBLISHED"
    elif batch.get("reviewQueueGenerated"):
        batch["status"] = "REVIEW_IN_PROGRESS"
    elif batch.get("mapDraftGenerated"):
        batch["status"] = "MAP_DRAFT_GENERATED"
    elif batch.get("readyForMapDraft"):
        batch["status"] = "READY_FOR_MAP_DRAFT"
    elif batch.get("readyForKeywordCalibration"):
        batch["status"] = "READY_FOR_KEYWORD_CALIBRATION"
    elif not batch.get("casesImported") and batch["knowledgeRequired"] and not batch.get("knowledgeImported"):
        batch["status"] = "WAITING_FOR_CASES_AND_KNOWLEDGE"
    elif not batch.get("casesImported"):
        batch["status"] = "WAITING_FOR_CASES"
    else:
        batch["status"] = "WAITING_FOR_KNOWLEDGE"


def parse_cases(file_name: str, content: bytes, max_cases: int = 10000) -> dict[str, Any]:
    if not (file_name or "").strip():
        raise ValueError("文件名不能为空")
    if not content:
        raise ValueError("用例文件为空")
    suffix = extension(file_name)
    if suffix in {"csv", "txt"}:
        text, risk = decode_text(content)
        matrix = list(csv.reader(io.StringIO(text)))
        risks = [risk] if risk else []
        return _parse_matrix(file_name, content, "csv", matrix, None, None, risks, max_cases)
    if suffix in {"xlsx", "xls"}:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), data_only=True)
        if not workbook.sheetnames:
            raise ValueError("Excel 文件没有工作表")
        sheet = workbook[workbook.sheetnames[0]]
        matrix = []
        for row in sheet.iter_rows(values_only=True):
            matrix.append(["" if cell is None else str(cell) for cell in row])
        return _parse_matrix(file_name, content, suffix, matrix, sheet.title, len(workbook.sheetnames), [], max_cases)
    if suffix == "xmind":
        parsed = parse_xmind_archive(content)
        rows = parsed["rows"]
        if not rows:
            raise ValueError("用例文件中没有可导入的用例")
        if len(rows) > max_cases:
            raise ValueError("单批次用例数不能超过 " + str(max_cases))
        return {
            "metadata": {
                "fileName": file_name, "format": parsed["format"], "size": len(content),
                "sha256": sha256_bytes(content), "rowCount": len(rows),
                "sheetName": None, "sheetCount": None,
            },
            "rows": rows,
            "risks": [{"code": "GENERATED_XMIND_CASE_IDS", "message": "已按 XMind 主题生成批次内临时用例 ID", "row": None}],
        }
    raise ValueError("暂不支持该用例文件格式：" + file_name)


def _parse_matrix(
    file_name: str, content: bytes, fmt: str, matrix: list[list[str]],
    sheet_name: str | None, sheet_count: int | None, risks: list[dict[str, Any]], max_cases: int,
) -> dict[str, Any]:
    first_content = next((index for index, row in enumerate(matrix) if any(_compact(cell) for cell in row)), -1)
    if first_content < 0:
        raise ValueError("用例文件为空")
    first_row = matrix[first_content]
    name_index = _find_header_index(first_row, CASE_NAME_KEYS)
    has_header = name_index >= 0
    if not has_header:
        mapping = {"id": -1, "name": 0, "step": 1, "expected": 2, "scene": 3, "feature": -1, "module": -1}
        risks.append({"code": "HEADER_NOT_DETECTED", "message": "未识别到表头，已按用例名称、步骤、预期、场景的默认列顺序解析", "row": first_content + 1})
        start_index = first_content
    else:
        mapping = {
            "id": _find_header_index(first_row, CASE_ID_KEYS),
            "name": name_index,
            "step": _find_header_index(first_row, STEP_KEYS),
            "expected": _find_header_index(first_row, EXPECTED_KEYS),
            "scene": _find_header_index(first_row, SCENE_KEYS),
            "feature": _find_header_index(first_row, FEATURE_KEYS),
            "module": _find_header_index(first_row, MODULE_KEYS),
        }
        _append_missing_column_risks(mapping, risks, first_content + 1)
        start_index = first_content + 1
    rows = []
    original_ids: set[str] = set()
    for matrix_index in range(start_index, len(matrix)):
        if len(rows) >= max_cases:
            raise ValueError("单批次用例数不能超过 " + str(max_cases))
        cells = matrix[matrix_index]
        if not any(_compact(cell) for cell in cells):
            continue
        source_row = matrix_index + 1
        case_name = _cell(cells, mapping["name"])
        if not case_name:
            risks.append({"code": "EMPTY_CASE_NAME", "message": "已跳过用例名称为空的行", "row": source_row})
            continue
        original_id = _cell(cells, mapping["id"])
        if not original_id:
            original_id = "IMPORT-ROW-" + str(source_row)
            risks.append({"code": "MISSING_ORIGINAL_CASE_ID", "message": "缺少原始用例 ID，已生成批次内临时 ID", "row": source_row})
        if original_id in original_ids:
            risks.append({"code": "DUPLICATE_ORIGINAL_CASE_ID", "message": "原始用例 ID 重复：" + original_id, "row": source_row})
        original_ids.add(original_id)
        rows.append({
            "originalCaseId": original_id,
            "caseName": case_name,
            "step": _cell(cells, mapping["step"]),
            "expected": _cell(cells, mapping["expected"]),
            "scene": _cell(cells, mapping["scene"]),
            "feature": _cell(cells, mapping["feature"]),
            "module": _cell(cells, mapping["module"]),
            "sourceRowNumber": source_row,
        })
    if not rows:
        raise ValueError("用例文件中没有可导入的用例")
    return {
        "metadata": {
            "fileName": file_name, "format": fmt, "size": len(content),
            "sha256": sha256_bytes(content), "rowCount": len(rows),
            "sheetName": sheet_name, "sheetCount": sheet_count,
        },
        "rows": rows,
        "risks": risks,
    }


def _cell(cells: list[str], index: int) -> str:
    if index < 0 or index >= len(cells):
        return ""
    return _compact(cells[index])


def _append_missing_column_risks(mapping: dict[str, int], risks: list[dict[str, Any]], header_row: int) -> None:
    if mapping["id"] < 0:
        risks.append({"code": "MISSING_CASE_ID_COLUMN", "message": "未识别原始用例 ID 列", "row": header_row})
    if mapping["step"] < 0:
        risks.append({"code": "MISSING_STEP_COLUMN", "message": "未识别测试步骤列", "row": header_row})
    if mapping["expected"] < 0:
        risks.append({"code": "MISSING_EXPECTED_COLUMN", "message": "未识别预期结果列", "row": header_row})
    if mapping["scene"] < 0:
        risks.append({"code": "MISSING_SCENE_COLUMN", "message": "未识别场景列", "row": header_row})
    if mapping["feature"] < 0:
        risks.append({"code": "MISSING_FEATURE_COLUMN", "message": "未识别功能点列", "row": header_row})


def parse_knowledge(file_name: str, content: bytes) -> dict[str, Any]:
    if not (file_name or "").strip():
        raise ValueError("文件名不能为空")
    if not content:
        raise ValueError("知识库文件为空")
    suffix = extension(file_name)
    if suffix in {"txt", "md", "markdown"}:
        text, risk = decode_text(content)
        if not text.strip():
            raise ValueError("知识库文件为空")
        return {
            "metadata": {
                "fileName": file_name, "format": suffix or "txt", "size": len(content),
                "sha256": sha256_bytes(content), "rowCount": None, "sheetName": None, "sheetCount": None,
            },
            "text": text,
            "lineCount": len(text.splitlines()),
            "risks": [risk] if risk else [],
        }
    if suffix in {"xlsx", "xls"}:
        return _parse_knowledge_workbook(file_name, content, suffix)
    if suffix == "xmind":
        parsed = parse_xmind_archive(content)
        text = parsed["knowledgeText"]
        if not text.strip():
            raise ValueError("知识库文件为空")
        return {
            "metadata": {
                "fileName": file_name, "format": parsed["format"], "size": len(content),
                "sha256": sha256_bytes(content), "rowCount": parsed["nodeCount"],
                "sheetName": None, "sheetCount": None,
            },
            "text": text,
            "lineCount": len(text.splitlines()),
            "risks": [],
        }
    raise ValueError("暂不支持该知识库文件格式：" + file_name)


def _parse_knowledge_workbook(file_name: str, content: bytes, suffix: str) -> dict[str, Any]:
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(content), data_only=True)
    if not workbook.sheetnames:
        raise ValueError("知识库 Excel 没有工作表")
    sheet = workbook[workbook.sheetnames[0]]
    knowledge_lines = []
    for row in sheet.iter_rows(values_only=True):
        cell_values = [_compact(cell) for cell in row if _compact(cell)]
        if cell_values:
            knowledge_lines.append("- " + "、".join(cell_values))
    if not knowledge_lines:
        raise ValueError("知识库 Excel 中没有可读取内容")
    text = "# 从 Excel 导入的知识草稿\n\n" + "\n".join(knowledge_lines)
    return {
        "metadata": {
            "fileName": file_name, "format": suffix, "size": len(content),
            "sha256": sha256_bytes(content), "rowCount": len(knowledge_lines),
            "sheetName": sheet.title, "sheetCount": len(workbook.sheetnames),
        },
        "text": text,
        "lineCount": len(text.splitlines()),
        "risks": [],
    }


_HTTP_METHODS = {"GET", "POST", "PUT", "PATCH", "DELETE"}
_API_FEATURE_HINTS = (
    ("findlist", ("列表查询", "列表", "页面测试")),
    ("count", ("统计", "角标")),
    ("enums", ("枚举",)),
    ("accept", ("接单", "认领")),
    ("transform", ("转派",)),
    ("urgentreview", ("催审",)),
    ("checkdesigncost", ("审核通过", "审核驳回", "驳回", "人工审核", "复审")),
    ("passdetail", ("明细",)),
    ("autodispatch", ("分派", "派单")),
    ("autosubmit", ("自动审核", "自动通过")),
    ("auditinfo", ("审核信息",)),
    ("wechat", ("企微",)),
    ("evaluate", ("评价",)),
    ("submitdesigncost", ("提交造价", "提交审核", "设计师提交")),
)
_API_CLASS_PREFIX = {
    "CostAuditApi": "/costAudit",
    "DesignCostApi": "/designCost",
    "TaskApi": "",
    "QuotationApi": "/quotation",
}
_IMPL_METHOD_PATHS = {
    "submitDesignCost": "/submitDesignCost",
    "autoSubmitDesignCost": "/autoSubmitDesignCost",
    "checkDesignCost": "/checkDesignCost",
    "passDetail": "/passDetail",
    "whetherAutoSubmit": "/whetherAutoSubmit",
}


def _strip_md_emphasis(text: str) -> str:
    return re.sub(r"[*_`]+", "", text or "").strip()


_JAVA_TYPE_SEGMENTS = {
    "date", "decimal", "string", "integer", "boolean", "long", "int", "map", "list", "object",
}


def _usable_api_path(path: str) -> bool:
    compact = (path or "").strip()
    if not compact.startswith("/") or len(compact) < 2:
        return False
    if compact in {"/", "/api", "/v1", "/v2"}:
        return False
    parts = [part for part in compact.strip("/").split("/") if part]
    if len(parts) >= 2 and all(len(part) <= 3 for part in parts):
        return False
    if parts and all(part.lower() in _JAVA_TYPE_SEGMENTS for part in parts):
        return False
    return True


def _api_label_from_text(text: str) -> str:
    compact = _compact(text).strip("`")
    if not compact:
        return ""
    matched = re.search(
        r"(GET|POST|PUT|PATCH|DELETE)\s*(?:[|｜,，]\s*)?`?(/[A-Za-z0-9_\-./{}]+)`?",
        compact,
        re.I,
    )
    if not matched:
        return ""
    path = matched.group(2).split("?")[0].rstrip(".,;")
    if not _usable_api_path(path):
        return ""
    return matched.group(1).upper() + " " + path


def _markdown_table_cells(line: str) -> list[str]:
    stripped = (line or "").strip()
    if not stripped.startswith("|"):
        return []
    cells = [_compact(cell).strip("`") for cell in stripped.strip("|").split("|")]
    if not any(cells):
        return []
    if all(re.fullmatch(r":?-{2,}:?", cell or "") for cell in cells):
        return []
    return cells


def _is_knowledge_table_header(cells: list[str]) -> bool:
    labels = {re.sub(r"\s+", "", cell) for cell in cells}
    header_marks = {
        "#", "方法", "路径", "入参", "出参", "场景", "说明", "代码证据",
        "功能点", "实现位置", "规则", "值", "字段", "含义", "接口",
    }
    return len(labels & header_marks) >= 2


def _table_column_index(headers: list[str], *names: str) -> int:
    normalized = [_normalize_header(item) for item in headers]
    wanted = {_normalize_header(name) for name in names}
    for index, label in enumerate(normalized):
        if label in wanted:
            return index
    return -1


def _api_label_from_cells(cells: list[str]) -> str:
    method = ""
    path = ""
    for cell in cells:
        if cell.upper() in _HTTP_METHODS:
            method = cell.upper()
            continue
        found = re.search(r"/[A-Za-z][A-Za-z0-9_\-./{}]*", cell)
        if not found:
            continue
        candidate = found.group(0).split("?")[0].rstrip(".,;")
        if _usable_api_path(candidate) and (candidate.count("/") >= 2 or method):
            path = candidate
    if path:
        return (method or "POST") + " " + path
    return _api_label_from_text(" ".join(cells))


def _table_feature_name(raw: str) -> str:
    cleaned = _strip_md_emphasis(raw)
    primary = re.sub(r"[（(][^）)]+[）)]", "", cleaned).strip() or cleaned
    return _compact(primary)


def _ingest_knowledge_table_row(
    cells: list[str],
    *,
    current_section: str,
    current_scene: str,
    current_feature: str,
    scenes: dict[str, dict[str, Any]],
    features: dict[str, dict[str, Any]],
    rules: dict[str, dict[str, Any]],
    apis: dict[str, dict[str, Any]],
    table_headers: list[str] | None = None,
) -> str:
    if _is_knowledge_table_header(cells):
        return current_feature
    api_label = _api_label_from_cells(cells)
    if api_label:
        feature_from_row = ""
        feature_index = _table_column_index(table_headers or [], "功能点", "功能", "feature")
        if feature_index >= 0 and feature_index < len(cells):
            feature_from_row = _table_feature_name(cells[feature_index])
            if _is_outline_label(feature_from_row) or _looks_like_api(feature_from_row):
                feature_from_row = ""
        if feature_from_row:
            _add_candidate(
                features,
                feature_from_row,
                "表格·功能点",
                80,
                related_scene=current_scene,
            )
        _add_candidate(
            apis,
            api_label,
            "表格·接口",
            88,
            related_scene=current_scene,
            related_feature=feature_from_row,
        )
        return feature_from_row or current_feature
    if current_section == "features":
        if cells and re.fullmatch(r"\d+", cells[0]) and len(cells) >= 2:
            feature_name = _table_feature_name(cells[1])
            impl_hint = " ".join(cells[2:])
        else:
            feature_name = _table_feature_name(cells[0] if cells else "")
            impl_hint = " ".join(cells[1:])
        if feature_name and not _is_outline_label(feature_name):
            _add_candidate(
                features,
                feature_name,
                "表格·功能点",
                88,
                related_scene=current_scene,
                impl_hint=impl_hint,
            )
            return feature_name
        return current_feature
    if current_section == "scenes" and cells:
        scene_name = _table_feature_name(cells[0])
        if scene_name and not _is_outline_label(scene_name):
            _add_candidate(scenes, scene_name, "表格·场景", 85)
        return current_feature
    if current_section == "rules" and cells:
        rule_name = _compact(cells[0])
        if rule_name and not _is_outline_label(rule_name) and len(rule_name) >= 4:
            _add_candidate(rules, rule_name, "表格·规则", 75, related_scene=current_scene)
    return current_feature


def _collect_impl_apis(source_text: str, apis: dict[str, dict[str, Any]]) -> None:
    for match in re.finditer(
        r"(CostAuditApi|DesignCostApi|TaskApi|QuotationApi)\s*[/#]([A-Za-z][A-Za-z0-9]+)",
        source_text or "",
    ):
        prefix = _API_CLASS_PREFIX.get(match.group(1), "")
        method_name = match.group(2)
        path = f"{prefix}/{method_name}" if prefix else f"/{method_name}"
        if _usable_api_path(path):
            _add_candidate(apis, "POST " + path, "实现·接口", 82)
    for method_name, path in _IMPL_METHOD_PATHS.items():
        if re.search(r"\b" + method_name + r"\b", source_text or ""):
            _add_candidate(apis, "POST " + path, "方法·接口", 78)


def _keep_real_apis(apis: dict[str, dict[str, Any]]) -> None:
    real: dict[str, dict[str, Any]] = {}
    for item in apis.values():
        label = _format_api(item.get("text") or "")
        if not label:
            continue
        payload = {**item, "text": label}
        current = real.get(label)
        if current is None:
            real[label] = payload
            continue
        if payload.get("relatedFeature") and not current.get("relatedFeature"):
            current["relatedFeature"] = payload["relatedFeature"]
        if payload.get("relatedScene") and not current.get("relatedScene"):
            current["relatedScene"] = payload["relatedScene"]
    apis.clear()
    apis.update(real)


def _link_apis_to_features(
    features: dict[str, dict[str, Any]],
    apis: dict[str, dict[str, Any]],
) -> None:
    feature_items = [item for item in features.values() if _usable_label(item.get("text") or "")]
    if not feature_items:
        return
    for api in apis.values():
        if _usable_label(api.get("relatedFeature") or ""):
            continue
        _method, path = _split_api(api.get("text") or "")
        last_key = re.sub(r"[^a-z0-9]", "", (path.split("?")[0].rstrip("/").split("/")[-1] if path else "").lower())
        path_key = re.sub(r"[^a-z0-9]", "", path.lower())
        best_text = ""
        best_score = 0
        best_scene = ""
        for feature in feature_items:
            text = feature.get("text") or ""
            impl = (feature.get("implHint") or "").replace("#", " ").replace("/", " ")
            blob = (text + " " + impl).lower()
            blob_compact = re.sub(r"[^a-z0-9\u4e00-\u9fff]", "", blob)
            score = 0
            if last_key and len(last_key) >= 5 and last_key in blob_compact:
                score += 24
            for english, chinese_hints in _API_FEATURE_HINTS:
                if english not in path_key:
                    continue
                if any(hint in text for hint in chinese_hints):
                    score += 12
            if score > best_score:
                best_score = score
                best_text = text
                best_scene = _usable_label(feature.get("relatedScene") or "")
        if best_text and best_score >= 12:
            api["relatedFeature"] = best_text
            if best_scene:
                api["relatedScene"] = best_scene


def extract_knowledge(text: str, source_version: str, process_node: str = "") -> dict[str, Any]:
    source_text = (text or "").strip()
    if not source_text:
        raise ValueError("知识文本不能为空")
    scenes: dict[str, dict[str, Any]] = {}
    features: dict[str, dict[str, Any]] = {}
    rules: dict[str, dict[str, Any]] = {}
    nodes: dict[str, dict[str, Any]] = {}
    apis: dict[str, dict[str, Any]] = {}
    has_sections = False
    current_scene = ""
    current_feature = ""
    current_section = ""
    table_headers: list[str] = []
    for line in source_text.splitlines():
        stripped = line.strip()
        heading = re.match(r"^(#{1,3})\s*(.+?)\s*$", stripped)
        outline_line = None if heading else re.match(
            r"^([一二三四五六七八九十百零〇0-9]+[、.．].+)$",
            stripped,
        )
        title = ""
        if heading:
            title = heading.group(2).strip()
        elif outline_line and _is_section_heading_line(outline_line.group(1)):
            title = outline_line.group(1).strip()
        if title:
            has_sections = True
            table_headers = []
            if "场景" in title:
                current_section = "scenes"
                scene_name = _business_label(title)
                if scene_name:
                    current_scene = scene_name
                    current_feature = ""
                    _add_candidate(scenes, scene_name, "章节标题·场景", 80)
            elif "功能" in title or "能力" in title:
                current_section = "features"
                feature_name = _business_label(title)
                if feature_name:
                    current_feature = feature_name
                    _add_candidate(features, feature_name, "章节标题·功能点", 80, related_scene=current_scene)
            elif "规则" in title or "校验" in title:
                current_section = "rules"
                rule_name = _business_label(title)
                if rule_name:
                    _add_candidate(rules, rule_name, "章节标题·规则", 75, related_scene=current_scene)
            elif "节点" in title:
                current_section = "nodes"
                node_name = _business_label(title)
                if node_name:
                    _add_candidate(nodes, node_name, "章节标题·节点", 75, related_scene=current_scene)
            elif "接口" in title or title.upper() == "API":
                current_section = "apis"
                api_name = _business_label(title)
                if api_name and not _is_outline_label(api_name) and _looks_like_api(api_name):
                    _add_candidate(apis, api_name, "章节标题·接口", 75, related_scene=current_scene, related_feature=current_feature)
            continue
        lead_scene = re.match(r"^(?:业务)?场景[:：]\s*(.+)$", stripped)
        if lead_scene:
            table_headers = []
            scene_name = _business_label(lead_scene.group(1))
            if scene_name:
                current_scene = scene_name
                current_feature = ""
                current_section = "scenes"
                _add_candidate(scenes, scene_name, "行首·场景", 85)
            continue
        lead_feature = re.match(r"^(?:功能点|功能|能力|模块)[:：]\s*(.+)$", stripped)
        if lead_feature:
            table_headers = []
            feature_name = _business_label(lead_feature.group(1)) or _compact(lead_feature.group(1))
            if feature_name and not _is_outline_label(feature_name):
                current_section = "features"
                current_feature = feature_name
                _add_candidate(features, feature_name, "行首·功能点", 85, related_scene=current_scene)
            continue
        lead_rule = re.match(r"^(?:规则|约束|校验)[:：]\s*(.+)$", stripped)
        if lead_rule:
            table_headers = []
            current_section = "rules"
            _add_candidate(rules, lead_rule.group(1), "行首·规则", 80, related_scene=current_scene)
            continue
        lead_node = re.match(r"^(?:流程节点|节点)[:：]\s*(.+)$", stripped)
        if lead_node:
            table_headers = []
            current_section = "nodes"
            _add_candidate(nodes, lead_node.group(1), "行首·节点", 80, related_scene=current_scene)
            continue
        lead_api = re.match(r"^(?:接口|API)[:：]\s*(.+)$", stripped, re.I)
        if lead_api:
            table_headers = []
            current_section = "apis"
            api_label = _api_label_from_text(lead_api.group(1)) or lead_api.group(1)
            _add_candidate(apis, api_label, "行首·接口", 85, related_scene=current_scene, related_feature=current_feature)
            continue
        cells = _markdown_table_cells(stripped)
        if cells:
            if _is_knowledge_table_header(cells):
                table_headers = cells
                if _table_column_index(table_headers, "路径", "接口", "方法") >= 0:
                    current_section = current_section or "apis"
                continue
            if current_section or _api_label_from_cells(cells):
                current_section = current_section or "apis"
                current_feature = _ingest_knowledge_table_row(
                    cells,
                    current_section=current_section,
                    current_scene=current_scene,
                    current_feature=current_feature,
                    scenes=scenes,
                    features=features,
                    rules=rules,
                    apis=apis,
                    table_headers=table_headers,
                )
                continue
        list_item = re.match(r"^[-*•、]\s*(.+)$", stripped)
        if list_item and current_section:
            item_text = _business_label(list_item.group(1)) or _compact(list_item.group(1))
            if item_text and not _is_outline_label(item_text):
                bags = {
                    "scenes": scenes,
                    "features": features,
                    "rules": rules,
                    "nodes": nodes,
                    "apis": apis,
                }
                _add_candidate(
                    bags[current_section],
                    item_text,
                    "列表·" + current_section,
                    80,
                    related_scene=current_scene,
                    related_feature=current_feature if current_section == "apis" else "",
                )
                if current_section == "features":
                    current_feature = item_text
            continue
    for match in re.finditer(r"(?:支持|提供|完成|实现|负责|用于)\s*([\u4e00-\u9fa5A-Za-z0-9]{2,16})", source_text):
        _add_candidate(features, match.group(1), "短文本·功能点", 65)
    for match in re.finditer(r"([\u4e00-\u9fa5A-Za-z0-9]{2,20}(?:必须|不可|不能|应当|禁止)[^\n。；;]{0,20})", source_text):
        _add_candidate(rules, match.group(1), "短文本·规则", 70)
    for match in re.finditer(
        r"(GET|POST|PUT|PATCH|DELETE)\s*(?:[|｜]\s*)?`?(\/[A-Za-z0-9_\-./{}]+)`?",
        source_text,
        re.I,
    ):
        path = match.group(2).split("?")[0].rstrip(".,;")
        if _usable_api_path(path):
            _add_candidate(apis, match.group(1).upper() + " " + path, "短文本·接口", 80)
    _collect_impl_apis(source_text, apis)
    _keep_real_apis(apis)
    _link_apis_to_features(features, apis)
    node = (process_node or "").strip()
    if node:
        _add_candidate(nodes, node, "任务·流程节点", 70)
    if not scenes and not features and not rules and not nodes and not apis:
        cleaned = _clean_item(source_text)
        if cleaned:
            _add_candidate(features, cleaned, "短文本·功能点", 65)
    if not scenes and not features and not rules and not nodes and not apis:
        raise ValueError("未从知识文本中抽取到有效候选，请上传知识文件或补充业务描述")
    return {
        "scenes": list(scenes.values())[:24],
        "features": list(features.values())[:30],
        "rules": list(rules.values())[:24],
        "nodes": list(nodes.values())[:24],
        "apis": list(apis.values())[:40],
        "mode": "structured+nl" if has_sections else "nl",
        "sourceVersion": source_version,
    }


def _clean_item(text: str) -> str:
    cleaned = re.sub(r"\s+", "", text)
    for word in STOP_WORDS:
        cleaned = cleaned.replace(word, "")
    return cleaned[:40]


def _add_candidate(
    bag: dict[str, dict[str, Any]],
    text: str,
    source: str,
    confidence: int,
    related_scene: str = "",
    related_feature: str = "",
    impl_hint: str = "",
) -> None:
    cleaned = _compact(text)
    if len(cleaned) < 2 or _is_outline_label(cleaned):
        return
    current = bag.get(cleaned)
    payload = {"text": cleaned, "source": source, "confidence": confidence}
    scene = _business_label(related_scene) if related_scene else ""
    if scene and _is_outline_label(scene):
        scene = ""
    feature = _usable_label(related_feature)
    hint = _compact(impl_hint)
    if scene:
        payload["relatedScene"] = scene
    elif current and current.get("relatedScene"):
        payload["relatedScene"] = current["relatedScene"]
    if feature:
        payload["relatedFeature"] = feature
    elif current and current.get("relatedFeature"):
        payload["relatedFeature"] = current["relatedFeature"]
    if hint:
        payload["implHint"] = hint
    elif current and current.get("implHint"):
        payload["implHint"] = current["implHint"]
    if current is None or confidence > current["confidence"]:
        bag[cleaned] = payload
    else:
        if scene and not current.get("relatedScene"):
            current["relatedScene"] = scene
        if feature and not current.get("relatedFeature"):
            current["relatedFeature"] = feature
        if hint and not current.get("implHint"):
            current["implHint"] = hint


def generate_map_draft(batch: dict[str, Any]) -> dict[str, Any]:
    case_import = batch.get("caseImport") or {}
    rows = case_import.get("rows") or []
    keywords = batch.get("keywordExtraction")
    if not rows:
        raise ValueError("请先导入历史用例")
    if keywords is None:
        raise ValueError("请先抽取并确认知识库关键字")
    if not batch.get("keywordsConfirmed"):
        raise ValueError("请先确认关键字")
    cleaned_cases = [_clean_case(row, keywords) for row in rows]
    process_node = str(batch.get("processNode") or batch.get("process_node") or "").strip()
    gaps = _build_gaps(cleaned_cases, keywords, process_node)
    stats = _draft_stats(cleaned_cases, gaps)
    return {
        "batchId": batch["id"],
        "cases": cleaned_cases,
        "gaps": gaps,
        "stats": stats,
        "generatedAt": _now(),
    }


def _clean_case(row: dict[str, Any], keywords: dict[str, Any]) -> dict[str, Any]:
    searchable = _normalize(" ".join([
        row.get("caseName") or "", row.get("step") or "", row.get("expected") or "",
        row.get("scene") or "", row.get("feature") or "", row.get("module") or "",
    ]))
    matched_scene = _best_match(searchable, _usable_keywords(keywords.get("scenes") or []))
    matched_feature = _best_match(searchable, _usable_keywords(keywords.get("features") or []))
    matched_node = _best_match(searchable, _usable_keywords(keywords.get("nodes") or []))
    historical_scene = _usable_label(row.get("scene"))
    historical_feature = _usable_label(row.get("feature"))
    identified_scene = _usable_label(matched_scene["text"] if matched_scene else "")
    identified_feature = _usable_label(matched_feature["text"] if matched_feature else "")
    knowledge_feature_names = {
        _usable_label(item.get("text"))
        for item in (keywords.get("features") or [])
        if _usable_label(item.get("text"))
    }
    if historical_feature in knowledge_feature_names:
        feature = historical_feature
    else:
        feature = identified_feature or historical_feature or "未识别功能点"
    scene = identified_scene or historical_scene or "未识别场景"
    matched_apis = _bind_knowledge_apis(
        scene=scene,
        feature=feature,
        searchable=searchable,
        keywords=keywords,
        allow_scene_fallback=False,
    )
    evidence = []
    confidence = 40
    if historical_scene:
        confidence += 15
        evidence.append("历史用例已有场景")
    if historical_feature:
        confidence += 20
        evidence.append("历史用例已有功能点")
    original_id = row.get("originalCaseId") or ""
    if not original_id.startswith("IMPORT-ROW-") and not original_id.startswith("XMIND-ROW-"):
        confidence += 5
        evidence.append("保留原始用例ID")
    if identified_scene:
        confidence += 5
        evidence.append("命中知识场景：" + identified_scene)
    elif historical_scene:
        evidence.append("知识未命中场景，已用历史用例场景兜底")
    if identified_feature:
        confidence += 10
        evidence.append("命中知识功能点：" + identified_feature)
    elif historical_feature:
        evidence.append("知识未命中功能点，已用历史用例功能点兜底")
    if matched_node:
        confidence += 5
        evidence.append("命中流程节点：" + matched_node["text"])
    if matched_apis:
        evidence.append("建议关联接口：" + _join_api_labels(matched_apis))
    confidence = min(confidence, 95)
    return {
        "originalCaseId": original_id,
        "caseName": row.get("caseName"),
        "scene": scene,
        "feature": feature,
        "flowNode": matched_node["text"] if matched_node else "",
        "api": _join_api_labels(matched_apis),
        "confidence": confidence,
        "mountAdvice": "自动挂载" if confidence >= 85 else ("待抽检" if confidence >= 70 else "人工确认"),
        "matchEvidence": evidence,
        "sourceRowNumber": row.get("sourceRowNumber"),
    }


def _best_match(searchable: str, candidates: list[dict[str, Any]]) -> dict[str, Any] | None:
    best = None
    for candidate in candidates:
        normalized = _normalize(candidate.get("text") or "")
        if len(normalized) < 2 or normalized not in searchable:
            continue
        if best is None or len(candidate.get("text") or "") > len(best.get("text") or ""):
            best = candidate
    return best


GAP_CASE_LIMIT = 180


def _usable_keywords(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    usable: list[dict[str, Any]] = []
    for item in items:
        text = item.get("text") or ""
        if _is_outline_label(text) or not _usable_label(text):
            continue
        usable.append(item)
    return usable


def _build_gaps(
    cleaned_cases: list[dict[str, Any]],
    keywords: dict[str, Any],
    process_node: str = "",
) -> list[dict[str, Any]]:
    corpus = [
        _normalize(" ".join([
            item.get("caseName") or "",
            item.get("step") or "",
            item.get("expected") or "",
            item.get("scene") or "",
            item.get("feature") or "",
            item.get("flowNode") or "",
        ]))
        for item in cleaned_cases
    ]
    gaps: list[dict[str, Any]] = []

    def append_cases(kind: str, source_item: dict[str, Any], reason: str) -> None:
        text = (source_item.get("text") or "").strip()
        if _skip_gap_text(text):
            return
        scene = _resolve_gap_scene(text, source_item, keywords, cleaned_cases, process_node)
        if kind == "接口缺口":
            feature_name = _api_feature_name(text)
            api_label = _format_api(text)
        else:
            feature_name = _resolve_gap_feature(text, keywords)
            api_label = _join_api_labels(
                _bind_knowledge_apis(
                    scene=scene,
                    feature=feature_name,
                    searchable=_normalize(" ".join([text, scene, feature_name])),
                    keywords=keywords,
                    allow_scene_fallback=True,
                )
            )
        variants = _gap_case_variants(kind, text if kind != "功能点缺口" else feature_name, scene, feature_name)
        if len(gaps) + len(variants) > GAP_CASE_LIMIT:
            return
        for variant in variants:
            gaps.append({
                "id": "GAP-" + str(len(gaps) + 1).zfill(3),
                "kind": kind,
                "scene": scene,
                "feature": feature_name,
                "api": api_label,
                "reason": reason,
                "caseName": variant["caseName"],
                "step": variant["step"],
                "expected": variant["expected"],
                "priority": variant["priority"],
            })

    real_features = _usable_keywords(keywords.get("features") or [])
    for feature in real_features:
        text = (feature.get("text") or "").strip()
        if _looks_like_api(text):
            continue
        if not _is_covered(text, corpus):
            append_cases("功能点缺口", feature, "知识库功能点未匹配历史用例")
    for rule in _usable_keywords(keywords.get("rules") or []):
        text = (rule.get("text") or "").strip()
        if not _is_covered(text, corpus):
            append_cases("规则缺口", rule, "知识库规则未匹配历史用例")
    for api in _usable_keywords(keywords.get("apis") or []):
        text = (api.get("text") or "").strip()
        if not _is_covered(text, corpus):
            append_cases("接口缺口", api, "知识库接口未匹配历史用例")
    return gaps


def _history_scene_fallback(cleaned_cases: list[dict[str, Any]]) -> str:
    counts: dict[str, int] = {}
    for item in cleaned_cases:
        scene = _usable_label(item.get("scene"))
        if not scene:
            continue
        counts[scene] = counts.get(scene, 0) + 1
    if not counts:
        return ""
    return max(counts.items(), key=lambda pair: (pair[1], len(pair[0])))[0]


def _resolve_gap_scene(
    text: str,
    source_item: dict[str, Any],
    keywords: dict[str, Any],
    cleaned_cases: list[dict[str, Any]],
    process_node: str = "",
) -> str:
    knowledge_scenes = _usable_keywords(keywords.get("scenes") or [])
    scene_names = {_usable_label(item.get("text")) for item in knowledge_scenes}
    scene_names.discard("")
    related = _usable_label(_business_label(source_item.get("relatedScene") or ""))
    node = _usable_label(process_node)
    if related and related != node and (not scene_names or related in scene_names):
        return related
    matched_scene = _best_match(text, knowledge_scenes)
    identified = _usable_label(matched_scene["text"] if matched_scene else "")
    if identified:
        return identified
    if len(scene_names) == 1:
        return next(iter(scene_names))
    historical = _history_scene_fallback(cleaned_cases)
    if historical:
        return historical
    return "未识别场景"


def _resolve_gap_feature(text: str, keywords: dict[str, Any]) -> str:
    label = _business_label(text) or (text or "").strip()
    if label and not _is_outline_label(label) and _usable_label(label):
        return label
    for feature in _usable_keywords(keywords.get("features") or []):
        name = (feature.get("text") or "").strip()
        if name:
            return name
    return _usable_label(label) or "未识别功能点"


def _numbered(*lines: str) -> str:
    return "\n".join(str(index) + ". " + line for index, line in enumerate(lines, 1))


def _gap_variant(case_name: str, priority: str, steps: tuple[str, ...], expected: tuple[str, ...]) -> dict[str, str]:
    return {
        "caseName": case_name,
        "priority": priority,
        "step": _numbered(*steps),
        "expected": _numbered(*expected),
    }


def _gap_case_variants(kind: str, text: str, scene: str, feature_name: str) -> list[dict[str, str]]:
    if kind == "规则缺口":
        return _functional_gap_variants("规则「" + text + "」", scene, text)
    if kind == "接口缺口":
        method, path = _split_api(text)
        return _api_gap_variants(method + " " + path, scene)
    return _functional_gap_variants("「" + feature_name + "」", scene, feature_name)


def _functional_gap_variants(subject: str, scene: str, detail: str) -> list[dict[str, str]]:
    return [
        _gap_variant(
            "验证" + subject + "正常业务主流程可完成且结果正确",
            "P1",
            (
                "进入场景「" + scene + "」，使用有权限账号打开" + subject + "入口",
                "按业务主路径准备合法、完整的前置数据",
                "执行该功能的核心操作并提交/保存",
                "核对页面展示、状态流转与关键计算结果",
            ),
            (
                "入口可见、可操作，页面加载完成",
                "主流程可走完，无阻塞级报错",
                "提交成功，核心数据与计算结果正确",
                "状态与下游展示与「" + detail + "」业务含义一致",
            ),
        ),
        _gap_variant(
            "验证" + subject + "必填项为空时拦截并提示",
            "P2",
            (
                "进入场景「" + scene + "」并打开" + subject,
                "将必填项留空或删除已填内容",
                "点击提交/保存",
                "查看校验提示与按钮、表单状态",
            ),
            (
                "必填项有明确标识",
                "空值无法作为有效数据提交",
                "提示指出具体必填字段，文案正确",
                "原数据不落库、不更新",
            ),
        ),
        _gap_variant(
            "验证" + subject + "在边界值与长度限制下处理正确",
            "P3",
            (
                "进入场景「" + scene + "」打开" + subject,
                "分别录入最小值、最大值、超长文本及临界长度",
                "提交并查看截断、拒绝或成功结果",
                "核对列表/详情中的展示与存储",
            ),
            (
                "合法边界值可提交且展示完整",
                "超限被拦截或按约定截断",
                "提示说明限制规则",
                "不出现计算溢出或页面错乱",
            ),
        ),
        _gap_variant(
            "验证" + subject + "对异常输入和特殊字符拦截或容错",
            "P2",
            (
                "进入场景「" + scene + "」打开" + subject,
                "录入空格、脚本字符、表情、全角符号或类型不符的值",
                "提交并观察前端校验与后端返回",
                "刷新页面确认未写入脏数据",
            ),
            (
                "非法值被拦截或按约定转义展示",
                "无脚本执行、无后端 5xx",
                "错误提示指向输入问题",
                "有效业务数据保持不变",
            ),
        ),
        _gap_variant(
            "验证" + subject + "重复提交不产生重复有效数据",
            "P2",
            (
                "进入场景「" + scene + "」完成一次合法提交",
                "在结果返回前或成功后立即再次点击提交",
                "查看是否生成第二条有效记录",
                "核对提示、按钮置灰与数据条数",
            ),
            (
                "首次提交成功或进入处理中",
                "重复点击不新增重复有效数据",
                "按钮处于加载/禁用或给出重复提示",
                "列表与详情仅保留一条有效结果",
            ),
        ),
        _gap_variant(
            "验证无权限角色无法访问或操作" + subject,
            "P2",
            (
                "使用无权限或只读角色登录",
                "尝试进入场景「" + scene + "」并打开" + subject,
                "尝试提交、编辑或删除",
                "改用有权限角色对照同一入口",
            ),
            (
                "无权限时入口隐藏、禁用或进入后提示无权限",
                "无法提交成功，接口返回权限错误",
                "不产生越权数据变更",
                "有权限角色可正常进入并操作",
            ),
        ),
        _gap_variant(
            "验证" + subject + "数据联动与状态流转正确",
            "P1",
            (
                "进入场景「" + scene + "」打开" + subject,
                "变更会驱动联动的字段、选项或上游状态",
                "观察关联字段、金额、按钮与下游状态是否同步",
                "提交后在列表、详情或下一节点核对状态",
            ),
            (
                "联动字段随源数据即时更新",
                "空值/零值/无数据时展示约定占位或禁用",
                "提交后状态与「" + detail + "」规则一致",
                "无滞后脏数据或错误状态残留",
            ),
        ),
        _gap_variant(
            "验证" + subject + "失败时错误提示文案正确且可理解",
            "P2",
            (
                "进入场景「" + scene + "」构造会失败的操作（缺参、规则不满足或依赖失败）",
                "执行提交/保存",
                "记录页面提示、字段红字与接口错误信息",
                "关闭提示后再次查看表单是否可继续修改",
            ),
            (
                "失败时给出明确中文提示，指向真实原因",
                "文案无堆栈、无英文代码裸奔（除非约定错误码）",
                "用户可据此修正后重试",
                "失败不改变原有效数据",
            ),
        ),
        _gap_variant(
            "验证" + subject + "按钮状态、页面交互与跳转符合规则",
            "P3",
            (
                "进入场景「" + scene + "」打开" + subject,
                "在未填完、校验失败、提交中、成功后分别观察按钮与加载态",
                "点击取消、返回、成功后的跳转链接或下一步",
                "使用浏览器后退确认页面与数据一致",
            ),
            (
                "不可用操作对应按钮禁用或不可点",
                "提交中防止重复点击，完成后恢复或跳转",
                "成功跳转到约定页面，失败停留并可改",
                "返回后数据与列表状态不错乱",
            ),
        ),
    ]


def _api_gap_variants(signature: str, scene: str) -> list[dict[str, str]]:
    return [
        _gap_variant(
            "验证" + signature + "合法请求走通主流程并返回成功",
            "P1",
            (
                "在场景「" + scene + "」准备有权限账号与完整合法入参",
                "调用" + signature + "一次",
                "核对响应码、关键字段与落库/下游结果",
                "在页面或查询接口确认主流程完成",
            ),
            (
                "鉴权通过，请求可发出",
                "返回约定成功码",
                "关键字段完整且与请求一致",
                "业务主流程完成，数据状态正确",
            ),
        ),
        _gap_variant(
            "验证" + signature + "必填参数为空时返回明确错误",
            "P2",
            (
                "构造缺少必填字段或字段为 null/空字符串的请求",
                "调用" + signature,
                "记录 HTTP 状态、业务码与错误信息",
                "查询确认未写入有效数据",
            ),
            (
                "请求被拒绝，返回 4xx 或业务失败码",
                "错误信息指出缺失字段",
                "不创建、不更新有效业务数据",
                "无 5xx",
            ),
        ),
        _gap_variant(
            "验证" + signature + "在边界值与长度限制下处理正确",
            "P3",
            (
                "分别构造最小合法值、最大合法值、超长字段请求",
                "调用" + signature,
                "核对成功与拒绝两类响应",
                "抽查存储长度与计算是否溢出",
            ),
            (
                "合法边界返回成功",
                "超限返回参数错误",
                "存储与展示不超过约定长度",
                "无计算溢出",
            ),
        ),
        _gap_variant(
            "验证" + signature + "异常类型与特殊字符被拒绝或安全转义",
            "P2",
            (
                "构造类型错误、脚本字符串、特殊符号入参",
                "调用" + signature,
                "查看错误码与返回体是否回显未转义脚本",
                "确认数据库无脏数据",
            ),
            (
                "非法入参失败，不落有效库",
                "返回体不执行、不回显可执行脚本",
                "错误指向参数问题",
                "无 5xx",
            ),
        ),
        _gap_variant(
            "验证" + signature + "重复调用不产生重复有效数据",
            "P2",
            (
                "用同一业务单据连续调用" + signature + "两次",
                "对比两次响应与数据条数",
                "检查是否幂等或第二次被拒绝",
                "核对列表仅一条有效结果（除非业务允许追加）",
            ),
            (
                "首次成功",
                "第二次幂等返回同一结果或明确拒绝重复",
                "不产生重复有效单据",
                "状态机不被重复推进两次",
            ),
        ),
        _gap_variant(
            "验证无权限调用" + signature + "被拒绝",
            "P2",
            (
                "使用过期 token、空 token 或无权限角色调用" + signature,
                "使用有权限账号对照调用一次",
                "核对未授权请求是否落库",
            ),
            (
                "无权限返回 401/403 或业务无权限码",
                "不产生越权写入",
                "有权限调用可成功",
            ),
        ),
        _gap_variant(
            "验证" + signature + "对 null/0/无数据的联动与状态处理正确",
            "P1",
            (
                "构造关联对象不存在、字段为 null、数值为 0 或列表为空的请求",
                "调用" + signature,
                "核对返回结构、默认值与页面/下游状态",
                "确认不会把空数据当成成功业务结果（除非约定）",
            ),
            (
                "空数据按约定返回空列表、默认值或业务失败",
                "不出现空指针式 5xx",
                "状态与金额因子为 null/0 时容错符合规则",
                "有数据时联动字段正确",
            ),
        ),
        _gap_variant(
            "验证" + signature + "失败时错误码与提示文案正确",
            "P2",
            (
                "分别构造缺参、无权限、规则不满足三类失败请求",
                "调用" + signature,
                "对比错误码、message 是否可区分原因",
            ),
            (
                "不同失败原因对应不同或可区分的错误信息",
                "文案可读，能指导修正",
                "失败不改变原有效数据",
            ),
        ),
        _gap_variant(
            "验证调用" + signature + "后前端按钮、加载与跳转状态正确",
            "P3",
            (
                "从场景「" + scene + "」页面触发会调用" + signature + "的按钮",
                "观察请求中、成功、失败时按钮与加载态",
                "成功后确认跳转或列表刷新；失败后停留可改",
            ),
            (
                "请求中按钮禁用或展示加载",
                "成功跳转或刷新到正确状态",
                "失败提示后可再次提交",
                "无重复提交导致的重复单据",
            ),
        ),
    ]


def _split_api(text: str) -> tuple[str, str]:
    cleaned = (text or "").strip().strip("`")
    matched = re.match(r"^(GET|POST|PUT|PATCH|DELETE)\s+(\S+)", cleaned, re.I)
    if matched:
        path = matched.group(2).strip("`,;，。")
        return matched.group(1).upper(), path
    return "POST", cleaned or "/unknown"


def _format_api(text: str) -> str:
    method, path = _split_api(text)
    if not path or path == "/unknown" or not path.startswith("/"):
        return ""
    if not _usable_api_path(path):
        return ""
    return method + " " + path.split("?")[0]


def _join_api_labels(labels: list[str]) -> str:
    unique: list[str] = []
    seen: set[str] = set()
    for label in labels:
        compact = (label or "").strip()
        if not compact or compact in seen:
            continue
        seen.add(compact)
        unique.append(compact)
    return "; ".join(unique)


def _split_api_labels(value: str | None) -> list[str]:
    if not value:
        return []
    return [part.strip() for part in re.split(r"[;\n]+", str(value)) if part.strip()]


def _bind_knowledge_apis(
    *,
    scene: str,
    feature: str,
    searchable: str,
    keywords: dict[str, Any],
    allow_scene_fallback: bool,
) -> list[str]:
    candidates = keywords.get("apis") or []
    text_hits = _match_apis(searchable, candidates)
    scene_name = _usable_label(scene)
    feature_name = _usable_label(feature)
    feature_norm = _normalize(feature_name)
    feature_hits: list[str] = []
    scene_hits: list[str] = []
    action_hits: list[str] = []
    haystack = _normalize(" ".join([searchable, feature_name, scene_name]))
    for candidate in candidates:
        label = _format_api(candidate.get("text") or "")
        if not label:
            continue
        related_scene = _usable_label(candidate.get("relatedScene") or "")
        related_feature = _usable_label(candidate.get("relatedFeature") or "")
        blob = _normalize(" ".join([candidate.get("text") or "", label, related_feature]))
        if feature_name and related_feature == feature_name:
            feature_hits.append(label)
            continue
        if feature_norm and len(feature_norm) >= 4 and feature_norm in blob:
            feature_hits.append(label)
            continue
        method, path = _split_api(candidate.get("text") or "")
        path_key = re.sub(r"[^a-z0-9]", "", path.lower())
        path_norm = _normalize(path.replace("/", " "))
        if feature_norm and len(feature_norm) >= 4 and path_norm and feature_norm in path_norm:
            feature_hits.append(label)
            continue
        for english, chinese_hints in _API_FEATURE_HINTS:
            if english not in path_key:
                continue
            if any(_normalize(hint) in haystack for hint in chinese_hints if len(hint) >= 2):
                action_hits.append(label)
                break
        if scene_name and related_scene == scene_name:
            scene_hits.append(label)
    ordered = _join_api_labels(text_hits + feature_hits + action_hits).split("; ") if (text_hits or feature_hits or action_hits) else []
    ordered = [item for item in ordered if item]
    if ordered:
        return ordered
    if allow_scene_fallback and 1 <= len(scene_hits) <= 3:
        return scene_hits
    return text_hits


def _match_apis(searchable: str, candidates: list[dict[str, Any]]) -> list[str]:
    matched: list[str] = []
    for candidate in candidates:
        text = (candidate.get("text") or "").strip()
        label = _format_api(text)
        if not label:
            continue
        tokens = {_normalize(text), _normalize(label)}
        method, path = _split_api(text)
        if path.startswith("/"):
            tokens.add(_normalize(path))
            tokens.add(_normalize(path.replace("/", "")))
        tokens = {token for token in tokens if len(token) >= 4}
        if not any(token in searchable for token in tokens):
            continue
        matched.append(label)
    return matched


def _skip_gap_text(text: str) -> bool:
    compact = (text or "").strip()
    if len(compact) < 2:
        return True
    if _is_outline_label(compact):
        return True
    return compact in {"规则", "校验", "约束", "功能", "功能点", "能力", "接口", "API", "api", "场景", "节点"}


def _looks_like_api(text: str) -> bool:
    return bool(re.match(r"^(GET|POST|PUT|PATCH|DELETE)\b", (text or "").strip(), re.I) or (text or "").startswith("/"))


def _api_feature_name(text: str) -> str:
    return re.sub(r"/+", " ", text).strip() or "未命名接口"


def _is_covered(text: str, corpus: list[str]) -> bool:
    normalized = _normalize(text)
    folded = normalized.replace("/", "")
    if len(folded) < 2:
        return True
    for blob in corpus:
        if not blob:
            continue
        blob_folded = blob.replace("/", "")
        if folded in blob_folded or blob_folded in folded:
            return True
    return False


def _draft_stats(cleaned_cases: list[dict[str, Any]], gaps: list[dict[str, Any]]) -> dict[str, int]:
    auto_mounted = sum(1 for item in cleaned_cases if item["confidence"] >= 85)
    review_required = sum(1 for item in cleaned_cases if 70 <= item["confidence"] < 85)
    manual_required = sum(1 for item in cleaned_cases if item["confidence"] < 70)
    return {
        "totalCases": len(cleaned_cases),
        "autoMounted": auto_mounted,
        "reviewRequired": review_required,
        "manualRequired": manual_required,
        "gapCount": len(gaps),
    }


def _normalize(value: str) -> str:
    return re.sub(r"\s+", "", (value or "").lower())


class ProduceService:
    def __init__(
        self,
        hierarchy_store: HierarchyStore,
        case_assets: CaseAssetStore,
        batch_file: Path,
        review_file: Path,
    ) -> None:
        self.hierarchy_store = hierarchy_store
        self.case_assets = case_assets
        self.batches = JsonStore(batch_file, "batches")
        self.reviews = JsonStore(review_file, "items")

    def create(self, body: dict[str, Any]) -> dict[str, Any]:
        mode = body.get("mode") or "standard"
        if mode != "standard":
            raise ValueError("当前仅开放标准模式")
        context = resolve_project_context(body)
        current_time = _now()
        batch = {
            "id": str(uuid4()),
            "domain": context["domain"],
            "system": context["system"],
            "moduleName": context["moduleName"],
            "processNode": context["processNode"],
            "ziroom_domain": context["domain"],
            "system_ref": context["system"],
            "module_name": context["moduleName"],
            "process_node": context["processNode"],
            "mode": "standard",
            "casesImported": False,
            "knowledgeImported": False,
            "keywordsConfirmed": False,
            "mapDraftGenerated": False,
            "reviewQueueGenerated": False,
            "publishedCount": 0,
            "caseImport": None,
            "knowledgeImport": None,
            "keywordExtraction": None,
            "mapDraft": None,
            "createdAt": current_time,
            "updatedAt": current_time,
        }
        refresh_gate(batch)
        self.batches.items[batch["id"]] = batch
        self.batches.save()
        return batch

    def find(self, batch_id: str) -> dict[str, Any] | None:
        return self.batches.items.get(batch_id)

    def get(self, batch_id: str) -> dict[str, Any]:
        batch = self.find(batch_id)
        if batch is None:
            raise ValueError("导入批次不存在：" + batch_id)
        return batch

    def import_cases(self, batch_id: str, file_name: str, content: bytes) -> dict[str, Any]:
        batch = self._editable(batch_id)
        parsed = parse_cases(file_name, content)
        batch["caseImport"] = parsed
        batch["casesImported"] = True
        batch["keywordsConfirmed"] = False
        batch["mapDraftGenerated"] = False
        batch["mapDraft"] = None
        batch["updatedAt"] = _now()
        refresh_gate(batch)
        self.batches.save()
        return parsed

    def import_knowledge(self, batch_id: str, file_name: str, content: bytes) -> dict[str, Any]:
        batch = self._editable(batch_id)
        parsed = parse_knowledge(file_name, content)
        batch["knowledgeImport"] = parsed
        batch["knowledgeImported"] = True
        batch["keywordExtraction"] = None
        batch["keywordsConfirmed"] = False
        batch["mapDraftGenerated"] = False
        batch["mapDraft"] = None
        batch["updatedAt"] = _now()
        refresh_gate(batch)
        self.batches.save()
        return parsed

    def extract_knowledge(self, batch_id: str) -> dict[str, Any]:
        batch = self._editable(batch_id)
        knowledge = batch.get("knowledgeImport") or {}
        text = knowledge.get("text") or ""
        if not text.strip():
            raise ValueError("请先上传或粘贴知识库内容")
        source_version = (knowledge.get("metadata") or {}).get("sha256") or sha256_bytes(text.encode("utf-8"))
        process_node = str(batch.get("processNode") or batch.get("process_node") or "").strip()
        extraction = extract_knowledge(text, source_version, process_node)
        batch["keywordExtraction"] = extraction
        batch["keywordsConfirmed"] = False
        batch["mapDraftGenerated"] = False
        batch["mapDraft"] = None
        batch["updatedAt"] = _now()
        refresh_gate(batch)
        self.batches.save()
        return extraction

    def import_knowledge_text(self, batch_id: str, source_name: str | None, text: str | None) -> dict[str, Any]:
        batch = self._editable(batch_id)
        normalized = (text or "").strip()
        if not normalized:
            raise ValueError("知识文本不能为空")
        existing = ((batch.get("knowledgeImport") or {}).get("text") or "").strip()
        combined = normalized if not existing else existing + "\n\n# 补充自然语言知识\n\n" + normalized
        content = combined.encode("utf-8")
        name = (source_name or "").strip() or "pasted-knowledge.md"
        if existing:
            previous_name = ((batch.get("knowledgeImport") or {}).get("metadata") or {}).get("fileName") or "knowledge"
            name = previous_name + " + " + name
        metadata = {
            "fileName": name,
            "format": "pasted-text" if not existing else "combined-knowledge",
            "size": len(content),
            "sha256": sha256_bytes(content),
            "rowCount": None, "sheetName": None, "sheetCount": None,
        }
        process_node = str(batch.get("processNode") or batch.get("process_node") or "").strip()
        extraction = extract_knowledge(combined, metadata["sha256"], process_node)
        batch["knowledgeImport"] = {"metadata": metadata, "text": combined, "lineCount": combined.count("\n") + 1, "risks": []}
        batch["knowledgeImported"] = True
        batch["keywordExtraction"] = extraction
        batch["keywordsConfirmed"] = False
        batch["mapDraftGenerated"] = False
        batch["mapDraft"] = None
        batch["updatedAt"] = _now()
        refresh_gate(batch)
        self.batches.save()
        return extraction

    def confirm_keywords(self, batch_id: str, body: dict[str, Any] | None) -> dict[str, Any]:
        batch = self._editable(batch_id)
        refresh_gate(batch)
        if not batch.get("readyForKeywordCalibration"):
            raise ValueError("标准模式需先完成历史用例和知识库导入")
        if batch.get("keywordExtraction") is None:
            raise ValueError("请先抽取知识库关键字")
        if body and any(body.get(key) for key in ("scenes", "features", "rules", "nodes", "apis")):
            current = batch["keywordExtraction"]
            batch["keywordExtraction"] = {
                **current,
                "scenes": _confirmed(body.get("scenes") or [], current.get("scenes") or []),
                "features": _confirmed(body.get("features") or [], current.get("features") or []),
                "rules": _confirmed(body.get("rules") or [], current.get("rules") or []),
                "nodes": _confirmed(body.get("nodes") or [], current.get("nodes") or []),
                "apis": _confirmed(body.get("apis") or [], current.get("apis") or []),
                "mode": "confirmed",
            }
        batch["keywordsConfirmed"] = True
        batch["updatedAt"] = _now()
        refresh_gate(batch)
        self.batches.save()
        return batch

    def generate_draft(self, batch_id: str) -> dict[str, Any]:
        batch = self._editable(batch_id)
        draft = generate_map_draft(batch)
        self._create_review_queue(batch, draft)
        batch["mapDraft"] = draft
        batch["mapDraftGenerated"] = True
        batch["reviewQueueGenerated"] = True
        batch["updatedAt"] = _now()
        refresh_gate(batch)
        self.batches.save()
        self.reviews.save()
        return draft

    def query_reviews(
        self, batch_id: str | None, system: str | None, status: str | None,
        kind: str | None, pool: str | None, keyword: str | None,
    ) -> dict[str, Any]:
        items = list(self.reviews.items.values())
        if batch_id:
            items = [item for item in items if item.get("batchId") == batch_id]
        if status:
            items = [item for item in items if item.get("status") == status]
        if kind:
            items = [item for item in items if item.get("kind") == kind]
        if keyword:
            needle = _normalize(keyword)
            items = [item for item in items if needle in _normalize(" ".join([
                str(item.get("caseName") or ""), str(item.get("scene") or ""),
                str(item.get("feature") or ""), str(item.get("originalCaseId") or ""),
                str(item.get("api") or ""),
            ]))]
        if pool == "A":
            items = [item for item in items if int(item.get("confidence") or 0) >= 85]
        elif pool == "B":
            items = [item for item in items if 70 <= int(item.get("confidence") or 0) < 85]
        elif pool == "C":
            items = [item for item in items if int(item.get("confidence") or 0) < 70]
        counts = {
            "total": len(items),
            "pending": sum(1 for item in items if item.get("status") == "待评审"),
            "confirmed": sum(1 for item in items if item.get("status") == "已确认"),
            "discarded": sum(1 for item in items if item.get("status") == "已废弃"),
            "published": sum(1 for item in items if item.get("status") == "已发布"),
        }
        return {"total": len(items), "counts": counts, "items": items}

    def update_review(self, review_id: str, body: dict[str, Any]) -> dict[str, Any]:
        item = self._review(review_id)
        for field_name, key in (
            ("caseName", "caseName"), ("scene", "scene"), ("feature", "feature"),
            ("step", "step"), ("expected", "expected"), ("priority", "priority"),
            ("flowNode", "flowNode"), ("api", "api"),
        ):
            if body.get(field_name) is not None:
                item[key] = str(body[field_name]).strip()
        item["updatedAt"] = _now()
        self.reviews.save()
        return item

    def confirm_review(self, review_id: str, body: dict[str, Any] | None) -> dict[str, Any]:
        item = self._pending(review_id)
        self._validate_confirm(item)
        item["status"] = "已确认"
        _add_op(item, "确认", _operator(body), (body or {}).get("comment") or "")
        item["updatedAt"] = _now()
        self.reviews.save()
        return item

    def discard_review(self, review_id: str, body: dict[str, Any] | None) -> dict[str, Any]:
        item = self._pending(review_id)
        item["status"] = "已废弃"
        _add_op(item, "废弃", _operator(body), (body or {}).get("comment") or "")
        item["updatedAt"] = _now()
        self.reviews.save()
        return item

    def confirm_batch(self, body: dict[str, Any]) -> dict[str, Any]:
        review_ids = body.get("reviewIds") or []
        succeeded = []
        failures = []
        for review_id in review_ids:
            try:
                self.confirm_review(review_id, body)
                succeeded.append(review_id)
            except ValueError as error:
                failures.append({"reviewId": review_id, "detail": str(error)})
        return {
            "requested": len(review_ids),
            "succeeded": len(succeeded),
            "failed": len(failures),
            "succeededIds": succeeded,
            "failures": failures,
        }

    def publish_batch(self, batch_id: str, body: dict[str, Any] | None) -> dict[str, Any]:
        batch = self.get(batch_id)
        confirmed = [
            item for item in self.reviews.items.values()
            if item.get("batchId") == batch_id and item.get("status") == "已确认"
        ]
        if not confirmed:
            raise ValueError("当前批次没有已确认且待发布的用例")
        published_at = _now()
        assets = [_to_official_asset(batch, item, published_at) for item in confirmed]
        self.case_assets.publish(assets)
        operator = _operator(body)
        for item in confirmed:
            item["status"] = "已发布"
            item["officialCaseId"] = _official_case_id(item["id"])
            _add_op(item, "发布", operator, (body or {}).get("comment") or "")
            item["updatedAt"] = published_at
        batch["publishedCount"] = int(batch.get("publishedCount") or 0) + len(assets)
        batch["updatedAt"] = published_at
        refresh_gate(batch)
        self.reviews.save()
        self.batches.save()
        return {
            "batchId": batch_id,
            "publishedCount": len(assets),
            "caseIds": [asset["id"] for asset in assets],
            "publishedAt": published_at,
        }

    def _editable(self, batch_id: str) -> dict[str, Any]:
        batch = self.get(batch_id)
        if batch.get("reviewQueueGenerated"):
            raise ValueError("当前批次已进入评审，不能重新导入或生成草稿")
        return batch

    def _review(self, review_id: str) -> dict[str, Any]:
        item = self.reviews.items.get(review_id)
        if item is None:
            raise ValueError("评审项不存在：" + review_id)
        return item

    def _pending(self, review_id: str) -> dict[str, Any]:
        item = self._review(review_id)
        if item.get("status") != "待评审":
            raise ValueError("仅待评审状态允许此操作，当前状态：" + str(item.get("status")))
        return item

    def _validate_confirm(self, item: dict[str, Any]) -> None:
        for field_name, label in (("caseName", "用例名称"), ("scene", "业务场景"), ("feature", "功能点")):
            if not (item.get(field_name) or "").strip():
                raise ValueError(label + "不能为空")
        for field_name, label in (("scene", "业务场景"), ("feature", "功能点")):
            value = item.get(field_name) or ""
            if "/" in value or "\n" in value or "\r" in value:
                raise ValueError(label + "不能包含斜杠或换行")
        if item.get("priority") not in {"P0", "P1", "P2", "P3", "P4"}:
            raise ValueError("优先级仅支持 P0、P1、P2、P3、P4")
        if item.get("kind") in {"知识缺口", "功能点缺口", "规则缺口", "接口缺口"}:
            if not (item.get("step") or "").strip():
                raise ValueError("缺口用例步骤不能为空")
            if not (item.get("expected") or "").strip():
                raise ValueError("缺口用例预期不能为空")

    def _create_review_queue(self, batch: dict[str, Any], draft: dict[str, Any]) -> None:
        source_rows = {(row.get("originalCaseId")): row for row in (batch.get("caseImport") or {}).get("rows") or []}
        keep_ids = []
        historical_index = 1
        current_time = _now()
        for cleaned in draft.get("cases") or []:
            review_id = _review_id(batch["id"], "H", historical_index)
            historical_index += 1
            source_row = source_rows.get(cleaned.get("originalCaseId")) or {}
            item = _base_review(review_id, batch["id"], "历史用例", current_time)
            item.update({
                "originalCaseId": cleaned.get("originalCaseId"),
                "caseName": cleaned.get("caseName"),
                "scene": cleaned.get("scene"),
                "feature": cleaned.get("feature"),
                "flowNode": cleaned.get("flowNode"),
                "api": cleaned.get("api") or "",
                "step": source_row.get("step") or "",
                "expected": source_row.get("expected") or "",
                "priority": "P0" if int(cleaned.get("confidence") or 0) >= 85 else "P1",
                "confidence": cleaned.get("confidence"),
                "matchEvidence": cleaned.get("matchEvidence") or [],
            })
            self.reviews.items[review_id] = item
            keep_ids.append(review_id)
        gap_index = 1
        for gap in draft.get("gaps") or []:
            review_id = _review_id(batch["id"], "G", gap_index)
            gap_index += 1
            gap_kind = str(gap.get("kind") or "知识缺口")
            item = _base_review(review_id, batch["id"], gap_kind, current_time)
            item.update({
                "originalCaseId": gap.get("id"),
                "caseName": gap.get("caseName") or (gap_kind + "：" + str(gap.get("feature") or "")),
                "scene": gap.get("scene"),
                "feature": gap.get("feature"),
                "flowNode": "",
                "api": gap.get("api") or "",
                "step": gap.get("step") or "",
                "expected": gap.get("expected") or "",
                "priority": gap.get("priority") or "P1",
                "confidence": 60,
                "matchEvidence": [gap.get("reason")],
            })
            self.reviews.items[review_id] = item
            keep_ids.append(review_id)
        for existing_id in list(self.reviews.items):
            existing = self.reviews.items[existing_id]
            if existing.get("batchId") == batch["id"] and existing_id not in keep_ids and existing.get("status") == "待评审":
                del self.reviews.items[existing_id]


def _confirmed(values: list[str], extracted: list[dict[str, Any]]) -> list[dict[str, Any]]:
    unique: list[str] = []
    for value in values:
        text = (value or "").strip()
        if text and text not in unique:
            unique.append(text)
    result = []
    for value in unique:
        original = next((item for item in extracted if item.get("text") == value), None)
        source = "人工确认" if not original else "人工确认·" + (original.get("source") or "")
        payload = {"text": value, "source": source, "confidence": 100}
        if original:
            if original.get("relatedScene"):
                payload["relatedScene"] = original["relatedScene"]
            if original.get("relatedFeature"):
                payload["relatedFeature"] = original["relatedFeature"]
            if original.get("implHint"):
                payload["implHint"] = original["implHint"]
        result.append(payload)
    return result


def _review_id(batch_id: str, kind: str, index: int) -> str:
    return "RV-" + batch_id[:8] + "-" + kind + "-" + str(index).zfill(3)


def _official_case_id(review_id: str) -> str:
    return "CA-" + review_id


def _operator(body: dict[str, Any] | None) -> str:
    operator = (body or {}).get("operator")
    return operator.strip() if operator else "当前用户"


def _add_op(item: dict[str, Any], action: str, operator: str, detail: str) -> None:
    operations = list(item.get("operations") or [])
    operations.append({"action": action, "operator": operator, "detail": (detail or "").strip(), "at": _now()})
    item["operations"] = operations


def _base_review(review_id: str, batch_id: str, kind: str, current_time: str) -> dict[str, Any]:
    return {
        "id": review_id,
        "batchId": batch_id,
        "kind": kind,
        "status": "待评审",
        "createdAt": current_time,
        "updatedAt": current_time,
        "operations": [{
            "action": "生成",
            "operator": "系统",
            "detail": "历史用例清洗进入评审" if kind == "历史用例" else "知识缺口进入评审",
            "at": current_time,
        }],
    }


def _to_official_asset(batch: dict[str, Any], item: dict[str, Any], published_at: str) -> dict[str, Any]:
    return {
        "id": _official_case_id(item["id"]),
        "name": item.get("caseName"),
        "priority": item.get("priority"),
        "testScenario": item.get("scene"),
        "featureName": item.get("feature"),
        "sceneName": item.get("scene"),
        "featureKey": join(batch["domain"], batch["system"], item.get("scene") or "", item.get("feature") or ""),
        "status": "已确认",
        "lifecycle": "已发布",
        "api": item.get("api") or "",
        "confidence": item.get("confidence"),
        "step": item.get("step"),
        "expected": item.get("expected"),
        "sourceBatchId": item.get("batchId"),
        "sourceReviewId": item["id"],
        "originalCaseId": item.get("originalCaseId"),
        "knowledgeSourceVersion": ((batch.get("keywordExtraction") or {}).get("sourceVersion") or ""),
        "sourceType": item.get("kind"),
        "flowNode": item.get("flowNode"),
        "moduleName": batch.get("moduleName") or "",
        "processNode": batch.get("processNode") or batch.get("process_node") or "",
        "createdAt": published_at,
        "updatedAt": published_at,
    }
